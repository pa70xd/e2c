"""Conversión determinista a JSON de lo que producción le pasa al repositorio
y de lo que devuelve en el correo de respuesta."""

from __future__ import annotations

import dataclasses
import datetime as dt
import math
import re
from decimal import Decimal
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import load_workbook

# Campos que producción llena con ``datetime.now()``: se enmascaran.
VOLATILE_KEYS = {"created_at", "updated_at"}

_DATE_IN_NAME = re.compile(r"\d{4}-\d{2}-\d{2}")


def to_jsonable(obj: Any, _key: str | None = None) -> Any:
    """Lleva cualquier valor a algo serializable y estable.

    NaN/NA se escriben como ``"<NaN>"``/``"<NA>"`` (no como ``null``) para que
    los goldens delaten los valores crudos que hoy llegan al repositorio.
    """
    if _key in VOLATILE_KEYS and obj is not None:
        return "<now>"
    if obj is None or isinstance(obj, (bool, str)):
        return obj
    if obj is pd.NA or obj is pd.NaT:
        return "<NA>"
    if isinstance(obj, float):
        if math.isnan(obj):
            return "<NaN>"
        if math.isinf(obj):
            return "<inf>" if obj > 0 else "<-inf>"
        return obj
    if isinstance(obj, int):
        return obj
    if isinstance(obj, Decimal):
        return str(obj)
    if isinstance(obj, (dt.datetime, pd.Timestamp)):
        return obj.isoformat()
    if isinstance(obj, dt.date):
        return obj.isoformat()
    if isinstance(obj, (bytes, bytearray)):
        return f"<bytes:{len(obj)}>"
    if hasattr(obj, "item") and callable(obj.item) and type(obj).__module__.startswith("numpy"):
        return to_jsonable(obj.item(), _key)
    if isinstance(obj, pd.DataFrame):
        return {
            "__dataframe__": {
                "columns": [str(c) for c in obj.columns],
                "rows": [
                    [to_jsonable(v) for v in row]
                    for row in obj.itertuples(index=False, name=None)
                ],
            }
        }
    if hasattr(obj, "model_dump"):  # pydantic v2
        return to_jsonable(obj.model_dump(), _key)
    if dataclasses.is_dataclass(obj) and not isinstance(obj, type):
        return to_jsonable(dataclasses.asdict(obj), _key)
    if isinstance(obj, dict):
        return {str(k): to_jsonable(v, str(k)) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [to_jsonable(v) for v in obj]
    if isinstance(obj, (set, frozenset)):
        return sorted((to_jsonable(v) for v in obj), key=repr)
    return f"<{type(obj).__name__}>"


def _cell(value: Any) -> Any:
    """Las plantillas traen fórmulas de relleno en cientos de renglones: se
    marcan sin copiar su texto."""
    if isinstance(value, str) and value.startswith("="):
        return "<formula>"
    return to_jsonable(value)


def _trim(rows: list[list[Any]]) -> list[list[Any]]:
    """Quita renglones sin datos (solo vacío/fórmula) y columnas vacías del final."""
    rows = [list(r) for r in rows]
    while rows and all(v is None or v == "<formula>" for v in rows[-1]):
        rows.pop()
    width = 0
    for r in rows:
        for i in range(len(r) - 1, -1, -1):
            if r[i] is not None and r[i] != "<formula>":
                width = max(width, i + 1)
                break
    return [r[:width] for r in rows]


def summarize_xlsx(name: str, content: bytes) -> dict[str, Any]:
    """Resumen estable de un Excel: nombre, hojas y filas (no los bytes).

    Solo la primera hoja lleva sus filas; las demás (catálogos de las
    plantillas, miles de renglones) quedan como título y tamaño."""
    stable_name = _DATE_IN_NAME.sub("<fecha>", name)
    summary: dict[str, Any] = {"name": stable_name}
    try:
        workbook = load_workbook(BytesIO(content), data_only=False)
    except Exception as exc:  # noqa: BLE001 - se documenta en el golden
        summary["error"] = f"{type(exc).__name__}"
        summary["size_bytes"] = len(content)
        return summary
    try:
        sheets = []
        for index, ws in enumerate(workbook.worksheets):
            entry: dict[str, Any] = {"title": ws.title, "state": ws.sheet_state}
            if index == 0:
                entry["rows"] = _trim([[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)])
            else:
                entry["max_row"] = ws.max_row
            sheets.append(entry)
        summary["sheets"] = sheets
    finally:
        workbook.close()
    return summary


def _squash(body: Any) -> Any:
    """El HTML de las plantillas viene con sangría: se colapsa el espacio en
    blanco para que los goldens sean legibles; el texto no cambia."""
    if isinstance(body, str):
        return re.sub(r"\s+", " ", body).strip()
    return to_jsonable(body)


def summarize_reply(draft: dict[str, Any]) -> dict[str, Any]:
    body = draft["body"]
    return {
        "in_reply_to": draft["in_reply_to"],
        "to": draft["to"],
        # Los mensajes de Heineken/Holanda a veces son dict (bug conocido):
        # se congela tal cual.
        "body": _squash(body),
        "attachments": [summarize_xlsx(n, c) for n, c in draft["attachments"]],
    }
