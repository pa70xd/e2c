"""Fixtures SINTÉTICOS de Imbera y Torrey (formato "Lay Out Repare").

    python -m tests.fixtures.imbera.build            # reescribe tests/fixtures/{imbera,torrey}/
    python -m tests.fixtures.imbera.build --check    # falla si lo versionado difiere

Mismas reglas que tests/fixtures/build_fixtures.py: todo inventado
(example.com, teléfonos 555-01xx). De los correos reales solo se tomó la
ESTRUCTURA del layout: hoja visible "RUTA", variante A (encabezado de una
fila) y variante B (fila de grupos combinados + columnas), hojas ocultas de
catálogo ("LinesByCategori", "ModelsByLines") y una fila basura.
"""

from __future__ import annotations

import datetime as dt
import json
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

FIXTURES_DIR = Path(__file__).resolve().parents[1]
_FIXED = dt.datetime(2026, 1, 1, 0, 0, 0)

HEADERS_A = [
    "Modelo", "Serie", "Falla Reportada & NOTAS ADICIONALES", "Fecha Factura",
    "Num. De Distribuidor", "N° de Cliente", "Municipio", "TIPO DE MOVIMIENTO", "CANAL",
]

# Variante B: (grupo, columna). Grupo vacío = sin grupo.
COLUMNAS_B = [
    ("", "Modelo del equipo a REPORTAR"),
    ("", "SERIE DEL EQUIPO"),
    ("", "Falla Reportada & NOTAS ADICIONALES"),
    ("", "Fecha Factura"),
    ("", "Número de distribuidor IMBERA / TORREY"),
    ("", "NUMERO DE CLIENTE QUE REPORTA EL EQUIPO"),
    ("DATOS DE DIRECCIÓN", "Calle"),
    ("DATOS DE DIRECCIÓN", "Número"),
    ("DATOS DE DIRECCIÓN", "Colonia"),
    ("DATOS DE DIRECCIÓN", "CP"),
    ("DATOS DE DIRECCIÓN", "Municipio"),
    ("DATOS DE DIRECCIÓN", "Estado"),
    ("CASO DE CONTAR CON LOS DATOS", "Latitud"),
    ("CASO DE CONTAR CON LOS DATOS", "Longitud"),
    ("NOMBRE DE QUIÉN RECIBIRÁ EL SERVICIO", "Nombre"),
    ("NOMBRE DE QUIÉN RECIBIRÁ EL SERVICIO", "Teléfono"),
]

CATALOGOS = {
    "LinesByCategori": [["Categoria", "Linea"], ["REFRIGERACION", "VERTICAL"], ["REFRIGERACION", "HORIZONTAL"]],
    "ModelsByLines": [["Linea", "Modelo"], ["VERTICAL", "MOD-V100"], ["HORIZONTAL", "MOD-H200"]],
}


def _tel(n: int) -> str:
    return f"55501{n:05d}"


def _libro() -> Workbook:
    wb = Workbook()
    wb.properties.creator = "e2c-tests"
    wb.properties.created = _FIXED
    wb.properties.modified = _FIXED
    return wb


def _catalogos_ocultos(wb: Workbook) -> None:
    for nombre, filas in CATALOGOS.items():
        ws = wb.create_sheet(nombre)
        for fila in filas:
            ws.append(fila)
        ws.sheet_state = "hidden"


def _guardar(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def layout_a(path: Path, filas: list[list], *, oculta_con_datos: list[list] | None = None) -> None:
    wb = _libro()
    ws = wb.active
    ws.title = "RUTA"
    ws.append(HEADERS_A)
    for fila in filas:
        ws.append(fila)
    _catalogos_ocultos(wb)
    if oculta_con_datos is not None:
        ws2 = wb.create_sheet("RUTA_2")
        ws2.append(HEADERS_A)
        for fila in oculta_con_datos:
            ws2.append(fila)
        ws2.sheet_state = "hidden"
    _guardar(wb, path)


def layout_b(path: Path, filas: list[list]) -> None:
    wb = _libro()
    ws = wb.active
    ws.title = "RUTA"
    ws.append(["LAY OUT REPARE"])  # título
    ws.append([g for g, _ in COLUMNAS_B])
    ws.append([c for _, c in COLUMNAS_B])
    # Celdas combinadas de los grupos (fila 2): columnas consecutivas con
    # el mismo grupo.
    i = 0
    while i < len(COLUMNAS_B):
        j = i
        while j + 1 < len(COLUMNAS_B) and COLUMNAS_B[j + 1][0] == COLUMNAS_B[i][0]:
            j += 1
        if COLUMNAS_B[i][0] and j > i:
            ws.merge_cells(start_row=2, start_column=i + 1, end_row=2, end_column=j + 1)
        i = j + 1
    for fila in filas:
        ws.append(fila)
    _catalogos_ocultos(wb)
    _guardar(wb, path)


def respuesta_os(path: Path, filas: list[list]) -> None:
    wb = _libro()
    ws = wb.active
    ws.title = "RUTA"
    ws.append(["Modelo", "Serie", "OS"])
    for fila in filas:
        ws.append(fila)
    _guardar(wb, path)


def _write_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


# --------------------------------------------------------------------------
# IMBERA
# --------------------------------------------------------------------------

FILA_BASURA_A = ["", "", "", "", "", "", "", "RUTA", "TRADICIONAL"]


def build_imbera(out: Path) -> None:
    d = out / "imbera"
    layout_a(d / "lay_out_repare_ruta.xlsx", [
        ["MOD-V100", "SER-IMB-0001", "No enfría", "2026-01-15", "D-1001", "C-2001", "MONTERREY", "RUTA", "TRADICIONAL"],
        FILA_BASURA_A,
        ["MOD-V100", "", "Ruido en compresor", "", "D-1001", "C-2002", "MONTERREY", "RUTA", "TRADICIONAL"],
        ["MOD-H200", "SER-IMB-0003", "Retiro de equipo", "", "D-1001", "C-2003", "MONTERREY", "RETIRO", "TRADICIONAL"],
        ["MOD-H200", "SER-IMB-0004", "Instalación nueva", "", "D-1001", "C-2004", "MONTERREY", "INSTALACION", "INSTITUCIONAL"],
    ])
    layout_a(d / "lay_out_repare_garantia.xlsx", [
        ["MOD-V100", "SER-IMB-0101", "Equipo nuevo dañado de fábrica", "2026-02-01", "D-1002", "C-2101", "SALTILLO", "RUTA", "MODERNO"],
        FILA_BASURA_A,
    ])
    layout_a(d / "lay_out_repare_reenvio.xlsx", [
        ["MOD-V100", "SER-IMB-0002", "Ruido en compresor", "", "D-1001", "C-2002", "MONTERREY", "RUTA", "TRADICIONAL"],
    ])
    respuesta_os(d / "respuesta_os.xlsx", [
        ["MOD-V100", "SER-IMB-0001", "4031000001"],
        ["MOD-H200", "SER-IMB-0004", ""],
    ])
    (d / "factura.pdf").write_bytes(b"%PDF-1.4\n% factura sintetica e2c-tests\n%%EOF\n")
    _write_json(d / "correos.json", {
        "init_ruta": {
            "subject": "ORDEN DE SERVICIO JUMEX -- MONTERREY -- REP-261026_1",
            "body": "Buen día.\nSe adjunta el layout con los equipos a revisar.\nSaludos,\nContacto Imbera Ejemplo",
        },
        "init_garantia": {
            "subject": "SOLICITUD DE GARANTIA OS123456 COMERCIALIZADORA EJEMPLO",
            "body": "Hola.\nEquipo nuevo dañado al recibirlo. Adjunto factura y layout.\nSaludos",
        },
        "seguimiento": {
            "subject": "Seguimiento reporte REP-201026_3",
            "body": "Buen día, ¿nos apoyan con el estatus del reporte?\nGracias",
        },
        "outbound_excel": {
            "subject": "RE: ORDEN DE SERVICIO JUMEX -- MONTERREY -- REP-261026_1",
            "body": "Hola. Adjunto las OS generadas.\nSaludos,\nOperador REPARE Ejemplo",
        },
        "outbound_texto": {
            "subject": "RE: ORDEN DE SERVICIO JUMEX -- MONTERREY -- REP-261026_1",
            "body": "Buen día.\nComparto orden generada: 4031000002\nSaludos,\nOperador REPARE Ejemplo",
        },
        "outbound_texto_series": {
            "subject": "RE: ORDEN DE SERVICIO JUMEX -- MONTERREY -- REP-261026_1",
            "body": "Equipo SER-IMB-0004 orden generada: 4031000005\nEquipo SER-IMB-0001 orden generada: 4031000006",
        },
        "outbound_espera_pieza": {
            "subject": "RE: ORDEN DE SERVICIO JUMEX -- MONTERREY -- REP-261026_1",
            "body": "Hola, el equipo se encuentra en espera de pieza.\nSaludos,\nOperador REPARE Ejemplo",
        },
        "inbound_reenvio": {
            "subject": "RE: ORDEN DE SERVICIO JUMEX -- MONTERREY -- REP-261026_1",
            "body": "Se corrige la serie del segundo equipo.\nSaludos",
        },
        "comandos": {"estatus": "bot.estatus", "ayuda": "bot.ayuda"},
    })


# --------------------------------------------------------------------------
# TORREY
# --------------------------------------------------------------------------

FILA_BASURA_B = [""] * len(COLUMNAS_B)


def _fila_b(n, modelo, serie, falla, nombre):
    return [
        modelo, serie, falla, "2026-02-10", "D-3001", f"C-40{n:02d}",
        "AV EJEMPLO", str(100 + n), "CENTRO", "64000", "MONTERREY", "NUEVO LEON",
        "25.6701", "-100.3101", nombre, _tel(n),
    ]


def build_torrey(out: Path) -> None:
    d = out / "torrey"
    layout_b(d / "lay_out_repare_torrey.xlsx", [
        _fila_b(1, "TOR-V300", "SER-TOR-0001", "No enfría", "RECIBE EJEMPLO UNO"),
        FILA_BASURA_B,
        _fila_b(2, "TOR-HL50", "SER-TOR-0002", "Máquina de hielo no produce", "RECIBE EJEMPLO DOS"),
    ])
    (d / "factura.pdf").write_bytes(b"%PDF-1.4\n% factura sintetica e2c-tests\n%%EOF\n")
    _write_json(d / "correos.json", {
        "init": {
            "subject": "Garantía equipos Torrey",
            "body": "Buen día.\nAdjunto Lay Out Repare y factura.\nSaludos,\nContacto Torrey Ejemplo",
        },
        "outbound_texto": {
            "subject": "RE: Garantía equipos Torrey",
            "body": "Comparto orden generada: 4032000001\nSaludos,\nOperador REPARE Ejemplo",
        },
        "comandos": {"ayuda": "bot.ayuda", "cancelar": "bot.cancelar"},
    })


# --------------------------------------------------------------------------

def build_all(out: Path = FIXTURES_DIR) -> None:
    out = Path(out)
    build_imbera(out)
    build_torrey(out)


def _dump_xlsx(path: Path) -> list:
    wb = load_workbook(path)
    try:
        return [
            [ws.title, ws.sheet_state, sorted(str(r) for r in ws.merged_cells.ranges),
             [list(row) for row in ws.iter_rows(values_only=True)]]
            for ws in wb.worksheets
        ]
    finally:
        wb.close()


def check(out: Path = FIXTURES_DIR) -> list[str]:
    diffs: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        build_all(Path(tmp))
        for built in sorted(Path(tmp).rglob("*")):
            if not built.is_file():
                continue
            rel = built.relative_to(tmp)
            current = out / rel
            if not current.exists():
                diffs.append(f"falta {rel}")
            elif built.suffix == ".xlsx":
                if _dump_xlsx(built) != _dump_xlsx(current):
                    diffs.append(f"difiere {rel}")
            elif built.read_bytes() != current.read_bytes():
                diffs.append(f"difiere {rel}")
    return diffs


if __name__ == "__main__":
    if "--check" in sys.argv:
        problems = check()
        for line in problems:
            print(line)
        sys.exit(1 if problems else 0)
    build_all()
    print(f"Fixtures de Imbera y Torrey regenerados en {FIXTURES_DIR}")
