"""Genera los Excel y los correos SINTÉTICOS de las pruebas.

    python tests/fixtures/build_fixtures.py            # reescribe tests/fixtures/<cliente>/
    python tests/fixtures/build_fixtures.py --check    # falla si lo versionado difiere del script

Reglas:
  * Todos los valores (nombres, teléfonos, correos, direcciones, series) son
    inventados. Los dominios son ``example.com`` y los teléfonos están en el
    rango ficticio 555-01xx.
  * NADA se copia de ``emails pruebas/``, de los ``*.json`` de la raíz ni de los
    CSV: solo se tomó de ahí la estructura (encabezados) que esperan los parsers.
  * Los Excel son mínimos: los renglones justos para ejercitar cada regla de
    negocio (aceptados, rechazados, duplicados).

Salida por cliente:
    tests/fixtures/<cliente>/*.xlsx
    tests/fixtures/<cliente>/correos.json   # asuntos y cuerpos de correo
"""

from __future__ import annotations

import datetime as dt
import json
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

FIXTURES_DIR = Path(__file__).resolve().parent
_FIXED = dt.datetime(2026, 1, 1, 0, 0, 0)


# --------------------------------------------------------------------------
# utilidades
# --------------------------------------------------------------------------

def _write_xlsx(path: Path, headers: list[str], rows: list[list], *, title_rows: list[list] | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    for fila in title_rows or []:
        ws.append(fila)
    ws.append(headers)
    for fila in rows:
        ws.append(fila)
    wb.properties.creator = "e2c-tests"
    wb.properties.created = _FIXED
    wb.properties.modified = _FIXED
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _write_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def _tel(n: int) -> str:
    return f"55501{n:05d}"


# --------------------------------------------------------------------------
# DANONE
# --------------------------------------------------------------------------

DANONE_HEADERS = [
    "ID CLIENTE", "FECHA DE SOLICITUD", "MODELO", "SERIE FABRICANTE", "SERIE DNN/BNF",
    "TIPO DE SERVICIO", "CANAL", "REGION", "CEDIS", "RAZON SOCIAL", "DIRECCION", "CP",
    "CIUDAD", "COLONIA", "LATITUD", "LONGITUD", "RUTA", "TELEFONO 1", "TELEFONO 2",
    "COMENTARIOS", "CBU",
]


def _danone_row(n, servicio, cbu="DANONE", lat="19.4301", lon="-99.1301", fab=None, dnn=None, cp="06000"):
    return [
        f"10000{n:02d}", "2026-03-02", "MODELO-A100",
        fab if fab is not None else f"FAB-D-{n:04d}",
        dnn if dnn is not None else f"DNN-D-{n:04d}",
        servicio, "TRADICIONAL", "CENTRO", "CEDIS NORTE",
        f"Tienda Ejemplo {n}", f"Calle Ficticia {100 + n}", cp, "Ciudad Ejemplo", "Colonia Ejemplo",
        lat, lon, f"RUTA-{n:02d}", _tel(n), _tel(n + 50), f"Comentario sintetico {n}", cbu,
    ]


def build_danone(out: Path) -> None:
    d = out / "danone"
    _write_xlsx(d / "solicitud_servicios.xlsx", DANONE_HEADERS, [
        _danone_row(1, "INSTALACIÓN"),
        _danone_row(2, "RETIRO", cbu="BONAFONT"),
        _danone_row(3, "REPARACION"),
        _danone_row(4, "INSTALACION", cbu="CBU-DESCONOCIDO"),   # rechazo: CBU
        _danone_row(5, "TRASLADO"),                              # rechazo: tipo (bug conocido, ver okf/riesgos/bugs.md)
        _danone_row(6, "RETIRO", lat="0", lon="0"),              # rechazo: coordenadas
        _danone_row(7, "INSTALACION", fab="", dnn=""),           # rechazo: sin serie
    ])
    _write_xlsx(d / "ampliacion_servicios.xlsx", DANONE_HEADERS, [
        _danone_row(11, "INSTALACION"),
        _danone_row(12, "RETIRO"),
    ])
    _write_xlsx(
        d / "respuesta_os.xlsx",
        ["ID CLIENTE", "SERIE FABRICANTE", "SERIE DNN/BNF", "TIPO DE SERVICIO", "OS"],
        [
            ["1000001", "FAB-D-0001", "DNN-D-0001", "INSTALACION", "4030000001"],    # ACCEPTED
            ["1000002", "FAB-D-0002", "DNN-D-0002", "RETIRO", "EN REVISION"],        # PENDING
            ["1000003", "FAB-D-0003", "DNN-D-0003", "REPARACION", ""],               # REJECTED (sin OS)
            ["1000004", "FAB-D-0004", "DNN-D-0004", "INSTALACION", "[RECHAZADO] Sin cobertura"],  # PENDING
        ],
    )
    _write_json(d / "correos.json", {
        "init": {
            "subject": "Solicitud de servicios REPARE 02/03/2026",
            "body": "Buen dia.\nFavor de realizar las instalaciones y retiros de los equipos del archivo adjunto.\nSaludos,\nContacto Danone Ejemplo",
        },
        "inbound": {
            "subject": "RE: Solicitud de servicios REPARE 02/03/2026",
            "body": "Buen dia.\nAdjunto dos equipos adicionales para la misma solicitud.\nSaludos,\nContacto Danone Ejemplo",
        },
        "outbound": {
            "subject": "RE: Solicitud de servicios REPARE 02/03/2026",
            "body": "Hola. Adjunto las OS generadas para los equipos de la solicitud.\nSaludos,\nOperador REPARE Ejemplo",
        },
        "bot": {
            "subject": "RE: Solicitud de servicios REPARE 02/03/2026",
            "body": "\U0001F916 Estimado usuario REPARE,\nEste es un mensaje automatico de ejemplo del bot.",
        },
        "comandos": {
            "ayuda": "bot.ayuda",
            "estatus": "bot.estatus",
            "cancelar": "bot.cancelar",
            "finalizar": "bot.finalizar",
            "global": "bot.global",
            "desconocido": "bot.inexistente",
        },
        "sin_proposito": {
            "subject": "Seguimiento a solicitud",
            "body": "Hola, solo quiero consultar el estatus de mi solicitud anterior.",
        },
    })


# --------------------------------------------------------------------------
# PENAFIEL
# --------------------------------------------------------------------------

PENAFIEL_HEADERS = [
    "NO CLIENTE DETALLISTA", "FECHA SOLICITUD", "MODELO", "SERIE", "SOLICITUD DE SERVICIO",
    "NOMBRE PUNTO DE VENTA", "CALLE", "ENTRE CALLE", "Y CALLE", "CP",
    "DELEGACION/MUNICIPIO/CIUDAD", "COLONIA", "LATITUD", "LONGITUD", "NUM. TEL", "FALLA", "CANAL",
]


def _penafiel_row(n, servicio, serie=None, id_cliente=None, lat="19.4402", lon="-99.1402"):
    return [
        id_cliente if id_cliente is not None else f"20000{n:02d}", "2026-03-02", "MODELO-P200",
        serie if serie is not None else f"PEN-S-{n:04d}", servicio,
        f"Abarrotes Ejemplo {n}", f"Avenida Inventada {200 + n}", "Calle Uno", "Calle Dos", "06100",
        "Ciudad Ejemplo", "Colonia Muestra", lat, lon, _tel(n + 100), f"Falla sintetica {n}", "TRADICIONAL",
    ]


def build_penafiel(out: Path) -> None:
    d = out / "penafiel"
    _write_xlsx(d / "solicitud_servicios.xlsx", PENAFIEL_HEADERS, [
        _penafiel_row(1, "INSTALACION"),
        _penafiel_row(2, "RETIRO"),
        _penafiel_row(3, "REPARACION"),
        _penafiel_row(4, "INSTALACION", serie=""),     # rechazo: serie vacía (conserva MODELO para no ser descartado por el reader)
        _penafiel_row(5, "TRASLADOS"),
        _penafiel_row(6, "RETIRO", id_cliente=""),   # rechazo: ID de cliente vacío
    ])
    _write_xlsx(d / "ampliacion_servicios.xlsx", PENAFIEL_HEADERS, [
        _penafiel_row(11, "INSTALACION"),
        _penafiel_row(12, "RETIRO"),
    ])
    _write_xlsx(
        d / "respuesta_os.xlsx",
        ["NO CLIENTE DETALLISTA", "MODELO", "SERIE", "SOLICITUD DE SERVICIO", "OS"],
        [
            ["2000001", "MODELO-P200", "PEN-S-0001", "INSTALACION", "4030000101"],
            ["2000002", "MODELO-P200", "PEN-S-0002", "RETIRO", "EN REVISION"],
            ["2000003", "MODELO-P200", "PEN-S-0003", "REPARACION", ""],
        ],
    )
    _write_json(d / "correos.json", {
        "init": {
            "subject": "Solicitud de servicio Penafiel 02/03/2026",
            "body": "Buen dia.\nSolicito la atencion de los siguientes servicios: instalaciones y retiros. Adjunto el formato.\nSaludos,\nContacto Penafiel Ejemplo",
        },
        "inbound": {
            "subject": "RE: Solicitud de servicio Penafiel 02/03/2026",
            "body": "Buen dia.\nAdjunto dos equipos adicionales.\nSaludos,\nContacto Penafiel Ejemplo",
        },
        "outbound": {
            "subject": "RE: Solicitud de servicio Penafiel 02/03/2026",
            "body": "Hola. Adjunto las OS generadas.\nSaludos,\nOperador REPARE Ejemplo",
        },
        "comandos": {
            "ayuda": "bot.help",
            "estatus": "bot.status",
            "cancelar": "bot.cancelar",
            "finalizar": "bot.cerrar",
            "global": "bot.global",
        },
    })


# --------------------------------------------------------------------------
# HOLANDA
# --------------------------------------------------------------------------

HOLANDA_HEADERS = [
    "TE", "ORIGEN", "FECHA", "CASO2", "SAP", "NOMBRE DEL CLIENTE", "NOMBRE NEGOCIO", "TELEFONO",
    "CALLE Y NUMERO", "COLONIA", "MUNICIPIO", "CP", "ESTADO", "LATITUD", "LONGITUD", "SEGMENTO",
    "CANAL", "AGRUPADOR", "REGION", "AGENCIA", "RUTA", "AGENCIA SATELITE", "TIPO DE QUEJA",
    "CONGELADOR", "MODELO", "SERIE", "OXXO", "CLASIFICADOR DE CASO", "SIN SAP / SIN CONGELADOR",
    "PROVEEDOR",
]


def _holanda_row(n, caso2, *, origen, fecha, tel, congelador=None, serie=None, te=None):
    return [
        te if te is not None else str(n), origen, fecha, caso2, f"SAP{n:05d}",
        f"Cliente Ficticio {n}", f"Neverita Ejemplo {n}", tel,
        f"Calle Inventada {300 + n}", "Colonia Muestra", "Municipio Ejemplo", "06200", "Estado Ejemplo",
        "19.4503", "-99.1503", "SEGMENTO X", "TRADICIONAL", "AGRUPADOR X", "REGION X", "AGENCIA X",
        f"RUTA-{n:02d}", "SATELITE X", "Falla de congelador sintetica",
        congelador if congelador is not None else f"CONG{n:05d}", "MODELO-H300",
        serie if serie is not None else f"HOL-S-{n:05d}", "NO", "Temas de congelador", "NO", "PROVEEDOR X",
    ]


def build_holanda(out: Path) -> None:
    d = out / "holanda"
    _write_xlsx(d / "Casos REPARE 02032026.xlsx", HOLANDA_HEADERS, [
        _holanda_row(1, "C100001H", origen="HELADIO", fecha=dt.datetime(2026, 3, 2), tel="5215550100001"),
        _holanda_row(2, "C200001G", origen="GDR", fecha="02/03/2026", tel="5550100002"),
        # mismo CONGELADOR que el renglón 1 -> duplicado intra-lote
        _holanda_row(3, "C100002H", origen="HELADIO", fecha=dt.datetime(2026, 3, 2), tel="5215550100003",
                     congelador="CONG00001", serie="HOL-S-00003"),
        _holanda_row(4, "C300001X", origen="", fecha="02/03/2026", tel="5550100004"),    # rechazo: origen
        _holanda_row(5, "C100005H", origen="HELADIO", fecha=dt.datetime(2026, 3, 2), tel="5215550100005",
                     serie=""),                                                           # rechazo: serie vacía
    ])
    _write_xlsx(
        d / "Casos REPARE OS.xlsx",
        ["CASO2", "OS"],
        [["C100001H", "4030000201"], ["C200001G", "4030000202"], ["C100002H", ""]],
    )
    _write_json(d / "correos.json", {
        "init": {
            "subject": "Casos REPARE ",
            "body": "Buen dia.\nSe adjuntan los casos de la semana.\nSaludos,\nContacto Holanda Ejemplo",
        },
        "outbound": {
            "subject": "RE: Casos REPARE ",
            "body": "Hola. Adjunto las OS de los casos.\nSaludos,\nOperador REPARE Ejemplo",
        },
        "comandos": {
            "estatus": "bot.estatus REPARE_000042",
            "cancelar": "bot.cancelar REPARE_000042",
            "finalizar": "bot.cerrar REPARE_000042",
            "sin_folio": "bot.estatus",
            "ayuda": "bot.ayuda",
            "global": "bot.global",
        },
        "fuera_de_asunto": {
            "subject": "Otro tema distinto",
            "body": "Este correo no pertenece al hilo de Holanda.",
        },
    })


# --------------------------------------------------------------------------
# HEINEKEN
# --------------------------------------------------------------------------

HEINEKEN_D2D_HEADERS = [
    "RAZON SOCIAL", "RESPONSABLE", "MODELO", "SERIE", "FALLA", "REPORTE IMAGEN", "CALLE Y NUMERO",
    "ENTRE CALLE", "Y CALLE", "COLONIA/POBLADO", "DELEGACION/MUNICIPIO/CIUDAD", "CP", "NUM. TEL",
    "HORARIO DE ATENCION", "NO CLIENTE DETALLISTA", "CEDIS/DISTRIBUIDORA", "OBSERVACIONES",
    "SOLICITUD DE SERVICIO", "ACTIVO", "ID REPORTE/TICKET/FOLIO/PAEEEM", "TECNICO", "OS", "SUCURSAL",
]


def _hnk_d2d_row(n, falla, *, serie, canal="CANAL MODERNO", cliente=None, servicio="MOVIMIENTO"):
    return [
        f"Comercial Ejemplo {n}", f"Responsable {n}", "MODELO-K400", serie, falla, "", f"Calle Falsa {400 + n}",
        "Calle A", "Calle B", "Colonia Muestra", "Municipio Ejemplo", "06300", _tel(n + 200),
        "09:00-18:00", cliente if cliente is not None else f"30000{n:02d}", canal, f"Observacion sintetica {n}",
        servicio, "123456789", f"TCK-{n:04d}", "", "", "ORIENTE",
    ]


HEINEKEN_RUTA_HEADERS = [
    "MODELO", "SERIE", "FALLA REPORTADA & NOTAS ADICIONALES", "HORA ATENCION", "NUMERO DE CASO/FOLIO",
    "N° DE CLIENTE", "NOMBRE PUNTO DE VENTA", "CALLE", "NUMERO", "ENTRE CALLE", "Y CALLE", "COLONIA",
    "MUNICIPIO", "CODIGO POSTAL", "LATITUD", "LONGITUD(-)", "TELEFONO Y/O FAX",
    "NOMBRE CONTACTO PUNTO DE VENTA", "TIPO DE MOVIMIENTO", "CANAL", "REGION", "SUCURSAL", "CLIENTE",
    "ORDEN", "AVISO", "OS",
]


def _hnk_ruta_row(n, *, movimiento="REPARACION MERCADO", cp="06400", serie=None):
    return [
        "MODELO-K400", serie if serie is not None else f"HNK-R-{n:04d}", f"No enfria sintetico {n}", "10:00", "",
        f"40000{n:02d}", f"Cantina Ejemplo {n}", f"Calle Imaginaria {500 + n}", str(n), "Calle A", "Calle B",
        "Colonia Muestra", "Municipio Ejemplo", cp, "19.4604", "-99.1604", _tel(n + 300),
        f"Contacto {n}", movimiento, "TRADICIONAL", "ORIENTE", "SUC 01", f"Cliente Ruta {n}",
        f"ORD-{n:05d}", f"AVI-{n:05d}", "",
    ]


def build_heineken(out: Path) -> None:
    d = out / "heineken"
    _write_xlsx(d / "precarga_d2d.xlsx", HEINEKEN_D2D_HEADERS, [
        _hnk_d2d_row(1, "ENTREGA DE ENF", serie="ASIGNAR SERIE GENERICA"),
        _hnk_d2d_row(2, "RETIRO DE ENF", serie="HNK-D-0002", canal="PUNTO DE VENTA"),
        _hnk_d2d_row(3, "RETIRO DE MOB", serie="ASIGNAR SERIE GENERICA", canal="CENTRO DE CONSUMO"),
        _hnk_d2d_row(4, "FALLA INVENTADA", serie="HNK-D-0004"),                 # rechazo: falla
        _hnk_d2d_row(5, "RETIRO DE ENF", serie="HNK-D-0005", cliente="N/A"),    # rechazo: sin cliente
    ])
    _write_xlsx(
        d / "respuesta_d2d.xlsx",
        ["MODELO", "SERIE", "NO CLIENTE DETALLISTA", "ACTIVO", "FALLA", "OS"],
        [
            ["MODELO-K400", "HNK-D-0002", "3000002", "123456789", "RETIRO DE ENF", "4030000301"],
            ["MODELO-K400", "HEI2026010100001", "3000001", "123456789", "ENTREGA DE ENF", "PENDIENTE DE CITA"],
        ],
    )
    _write_xlsx(
        d / "mantenimiento_ferguez.xlsx", HEINEKEN_RUTA_HEADERS,
        [
            _hnk_ruta_row(1),
            _hnk_ruta_row(2),
            _hnk_ruta_row(3, movimiento="INSTALACION"),     # rechazo: tipo de movimiento
            _hnk_ruta_row(4, cp=""),                        # rechazo: sin CP
        ],
        title_rows=[["Mantenimiento Ferguez - formato sintetico"]],
    )
    _write_xlsx(
        d / "respuesta_ruta.xlsx",
        ["MODELO", "SERIE", "ORDEN", "OS"],
        [
            ["MODELO-K400", "HNK-R-0001", "ORD-00001", "4030000401"],
            ["MODELO-K400", "HNK-R-0002", "ORD-00002", ""],
        ],
    )
    _write_json(d / "correos.json", {
        "d2d_init": {
            "subject": "Precarga D2D 02/03/2026",
            "body": "Buen dia.\nSe adjunta la precarga de movimientos del dia.\nSaludos,\nContacto Heineken Ejemplo",
        },
        "d2d_reenvio": {
            "subject": "RE: Precarga D2D 02/03/2026",
            "body": "Se reenvia la precarga con correcciones.\nSaludos,\nContacto Heineken Ejemplo",
        },
        "d2d_outbound": {
            "subject": "RE: Precarga D2D 02/03/2026",
            "body": "Hola. Adjunto las OS de la precarga.\nSaludos,\nOperador REPARE Ejemplo",
        },
        "ruta_init": {
            "subject": "Mantenimiento Ferguez",
            "body": "Buen dia.\nAdjunto los equipos para reparacion.\nSaludos,\nContacto Heineken Ejemplo",
        },
        "ruta_outbound": {
            "subject": "RE: Mantenimiento Ferguez",
            "body": "Hola. Adjunto las OS de mantenimiento.\nSaludos,\nOperador REPARE Ejemplo",
        },
        "comandos": {
            "estatus_d2d": "bot.estatus",
            "cancelar_d2d": "bot.cancelar",
            "estatus_ruta": "bot.estatus REPARE_000042",
            "ayuda": "bot.help",
            "global": "bot.global",
        },
    })


# --------------------------------------------------------------------------
# main
# --------------------------------------------------------------------------

def build_all(out: Path = FIXTURES_DIR) -> None:
    out = Path(out)
    build_danone(out)
    build_penafiel(out)
    build_holanda(out)
    build_heineken(out)


def _dump_xlsx(path: Path) -> list:
    wb = load_workbook(path)
    try:
        return [[list(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    finally:
        wb.close()


def check(out: Path = FIXTURES_DIR) -> list[str]:
    """Reconstruye en un directorio temporal y compara el CONTENIDO (los zip
    de Excel no son byte a byte estables). Regresa los archivos que difieren."""
    diffs: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        build_all(Path(tmp))
        for built in sorted(Path(tmp).rglob("*")):
            if not built.is_file():
                continue
            rel = built.relative_to(tmp)
            current = out / rel
            if not current.exists():
                diffs.append(f"falta {rel}")
            elif built.suffix == ".xlsx":
                if _dump_xlsx(built) != _dump_xlsx(current):
                    diffs.append(f"difiere {rel}")
            elif built.read_bytes() != current.read_bytes():
                diffs.append(f"difiere {rel}")
    return diffs


if __name__ == "__main__":
    if "--check" in sys.argv:
        problems = check()
        for line in problems:
            print(line)
        sys.exit(1 if problems else 0)
    build_all()
    print(f"Fixtures regenerados en {FIXTURES_DIR}")
