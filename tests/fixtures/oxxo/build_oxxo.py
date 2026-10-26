"""Genera los fixtures SINTÉTICOS de Oxxo (correos de plantilla y Excel).

    python tests/fixtures/oxxo/build_oxxo.py            # reescribe tests/fixtures/oxxo/
    python tests/fixtures/oxxo/build_oxxo.py --check    # sale 1 si lo versionado difiere

Todo está inventado: nombres, teléfonos (555-01xx), direcciones, CR, tickets y
folios; correos en example.com. Solo se tomó de los correos reales la FORMA de
las plantillas (etiquetas y orden). Los ``.xlsm`` son libros de openpyxl con
extensión .xlsm (sin macros): para los lectores es lo mismo.
"""

from __future__ import annotations

import datetime as dt
import json
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

OUT = Path(__file__).resolve().parent
_FIXED = dt.datetime(2026, 1, 1, 0, 0, 0)

EXCEL_INIT = "Solicitud OS Oxxo Tienda Ejemplo.xlsm"
EXCEL_REPLY = "OS Cuartos Frios Ejemplo.xlsm"

ENCABEZADOS = [
    "FOLIO", "ID CLIENTE", "FECHA DE SOLICITUD", "MODELO", "SERIE", "TIPO DE SERVICIO",
    "FALLA", "REGION", "PLAZA", "RAZON SOCIAL", "DIRECCION", "CP", "MUNICIPIO", "COLONIA",
    "TELEFONO",
]


def _fila(n: int, *, folio: str | None, serie: str, modelo: str, falla: str):
    return [
        folio, f"50EJ{n}", "02/03/2026", modelo, serie, "GARANTIA", falla, "NORTE",
        "PLAZA EJEMPLO", f"OXXO TIENDA EJEMPLO {n}", f"Calle Ficticia {100 + n}", "06000",
        "Municipio Ejemplo", "Colonia Ejemplo", f"55501{n:05d}",
    ]


def _xlsm(path: Path, filas: list[list], *, titulo: list[list] | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Lay Out Repare"
    for renglon in titulo or []:
        ws.append(renglon)
    ws.append(ENCABEZADOS)
    for fila in filas:
        ws.append(fila)
    wb.properties.creator = "e2c-tests"
    wb.properties.created = _FIXED
    wb.properties.modified = _FIXED
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


MANTENIMIENTO_CUERPO = """Buen día,
La tienda OXXO 50EJE NORTE TIENDA EJEMPLO UNO, ubicada en Calle Ficticia 100, Colonia Ejemplo, Municipio Ejemplo, nos reportó lo siguiente
Motivo: El equipo no enfría
Representante Propietario: Persona Ejemplo Uno
Prestador de Servicio: REPARE - Norte
Categoria: Refrigeración
Falla especifica: Cuarto frío sin temperatura. El número de ticket es: {ticket}. Prioridad: ALTA

Saludos cordiales,
Sistema de tickets Ejemplo"""

SEGUIMIENTO_CUERPO = """Fecha de generación de folio: 02/03/2026 09:30
Descripción Servicio (Plaza): NORTE
CR Tienda: 50EJE
Nombre de quien reporta: Persona Ejemplo Dos
Teléfono de contacto: 55 5501 0002
Correo Electrónico: persona.dos@example.com
Equipo: Cuarto frío
Marca/Fabricante del Equipo: Marca Ejemplo
Modelo: MOD-EJ-01
Numero Serie: SERIE-EJ-0001
Diagnostico y/o Falla: Compresor no arranca
Direccion de Envío de Garantía: Calle Ficticia 200, Colonia Ejemplo
Folio Referencia Mantenimiento: 10000001
Mensaje de Seguimiento: se cumplió el tiempo y no se ha atendido"""


def correos() -> dict:
    return {
        "mantenimiento": {
            "subject": "Abierto-Asignado ----> 10000001",
            "body": MANTENIMIENTO_CUERPO.format(ticket="10000001"),
        },
        "mantenimiento_sin_ticket": {
            "subject": "Abierto-Asignado ----> ",
            "body": MANTENIMIENTO_CUERPO.format(ticket="(pendiente)"),
        },
        "seguimiento": {
            "subject": "Seguimiento a Folio: GAR-0001",
            "body": SEGUIMIENTO_CUERPO,
        },
        "seguimiento_sin_folio": {
            "subject": "Seguimiento a Folio:",
            "body": SEGUIMIENTO_CUERPO,
        },
        "excel_init": {
            "subject": "Solicitud de Orden de Servicio Oxxo Tienda Ejemplo",
            "body": "Buen día. Adjunto el layout con las OS de garantía de la tienda.\nSaludos,\nPersona Ejemplo Tres",
        },
        "excel_reply": {
            "subject": "RE: Solicitud de generación de OS para cliente Ejemplo - OXO",
            "body": "Hola. Les comparto las 4 OS de garantía en el layout adjunto.\nSaludos,\nPersona Ejemplo Tres",
        },
        "excel_reply_conteo": {
            "subject": "RE: Solicitud de generación de OS para cliente Ejemplo - OXO",
            "body": "Hola. Les comparto las 5 OS de garantía en el layout adjunto.\nSaludos,\nPersona Ejemplo Tres",
        },
        "comandos": {
            "estatus": "bot.estatus REPARE_000042",
            "ayuda": "bot.ayuda",
            "sin_folio": "bot.estatus",
        },
        "fuera_de_plantilla": {
            "subject": "Cerrado ----> 10000001",
            "body": "Este correo no pertenece a ninguna plantilla de Oxxo.",
        },
    }


def build_all(out: Path = OUT) -> None:
    out.mkdir(parents=True, exist_ok=True)
    (out / "correos.json").write_text(
        json.dumps(correos(), indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    _xlsm(out / EXCEL_INIT, [
        _fila(1, folio="OX-0001", serie="SERIE-EJ-0101", modelo="MOD-EJ-01", falla="No enfría"),
        _fila(2, folio="OX-0002", serie="SERIE-EJ-0102", modelo="MOD-EJ-01", falla="Gotea"),
        _fila(3, folio=None, serie="SERIE-EJ-0103", modelo="MOD-EJ-02", falla="Ruido en compresor"),
    ])
    # Línea de cuartos fríos de Imbera: dos renglones de título antes del encabezado.
    _xlsm(
        out / EXCEL_REPLY,
        [_fila(n, folio=f"OX-01{n:02d}", serie=f"SERIE-CF-{n:04d}", modelo="CUARTO-FRIO-EJ", falla="Temperatura alta")
         for n in range(1, 5)],
        titulo=[["Lay Out Repare - Cuartos fríos"], []],
    )


def _dump(path: Path):
    wb = load_workbook(path, data_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()


def check(out: Path = OUT) -> list[str]:
    diffs: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        build_all(Path(tmp))
        for built in sorted(Path(tmp).rglob("*")):
            if not built.is_file():
                continue
            actual = out / built.relative_to(tmp)
            if not actual.exists():
                diffs.append(f"falta {built.name}")
            elif built.suffix == ".xlsm":
                if _dump(built) != _dump(actual):
                    diffs.append(f"difiere {built.name}")
            elif built.read_bytes() != actual.read_bytes():
                diffs.append(f"difiere {built.name}")
    return diffs


if __name__ == "__main__":
    if "--check" in sys.argv:
        problemas = check()
        for linea in problemas:
            print(linea)
        sys.exit(1 if problemas else 0)
    build_all()
    print(f"Fixtures de Oxxo regenerados en {OUT}")
