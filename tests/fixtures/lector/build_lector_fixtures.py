"""Genera los adjuntos SINTÉTICOS de tests/unit_core/test_lector.py.

    python tests/fixtures/lector/build_lector_fixtures.py

Todos los valores son inventados. Los formatos imitan la ESTRUCTURA de los
correos reales (encabezado agrupado, hojas ocultas de datos y de catálogo,
encabezados con saltos de línea, filas basura), no su contenido.
"""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

DIR = Path(__file__).resolve().parent

# --- vocabularios y claves que usan las pruebas -----------------------------
VOCAB_KOF = [
    "MODELO", "No SERIE", "FALLA", "DELEGACION/MUNICIPIO/CIUDAD",
    "No CLIENTE DETALLISTA- CONCATENADO", "ID DISTRIBUIDOR", "ID DETALLISTA",
    "CEDIS/DISTRIBUIDORA/ SUCURSAL", "OBSERVACIONES", "SOLICITUD DE SERVICIO",
    "ACTIVO/PLACA", "OS REPARE",
]
CLAVES_KOF = ["No SERIE", "MODELO", "ACTIVO/PLACA"]
VOCAB_RUTA = ["FOLIO", "No SERIE", "MODELO", "CALLE", "COLONIA", "MUNICIPIO", "TELEFONO"]
CLAVES_RUTA = ["FOLIO", "No SERIE", "MODELO"]

ENCABEZADO_KOF = [
    "MODELO", "No SERIE", "FALLA", "DELEGACION/\nMUNICIPIO/\nCIUDAD",
    "No CLIENTE DETALLISTA- CONCATENADO", "ID DISTRIBUIDOR", "ID DETALLISTA",
    "CEDIS/DISTRIBUIDORA/\nSUCURSAL", "OBSERVACIONES", "SOLICITUD DE SERVICIO",
    "ACTIVO/PLACA", "OS REPARE",
]


def _fila_kof(n: int) -> list:
    return [
        f"MOD-{n}", f"SER{n:05d}", "No enfría", "Ciudad Ejemplo", f"D{n}-R{n}",
        f"DIST{n}", f"DET{n}", "Cedis Ejemplo", None, f"SS{n}", f"PL{n}", None,
    ]


def _bytes(wb: Workbook) -> bytes:
    wb.properties.creator = "fixture"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def kof_xlsx() -> bytes:
    """Hoja visible REPORTE + 7 ocultas con el mismo encabezado + 1 catálogo oculto."""
    wb = Workbook()
    visible = wb.active
    visible.title = "REPORTE"
    visible.append(ENCABEZADO_KOF)
    visible.append(_fila_kof(1))
    for k in range(1, 8):
        h = wb.create_sheet(f"DIA{k}")
        h.sheet_state = "hidden"
        h.append(ENCABEZADO_KOF)
        h.append(_fila_kof(10 * k))
        h.append(_fila_kof(10 * k + 1))
        h.append(["solo una celda"] + [None] * 11)  # basura
    cat = wb.create_sheet("CATALOGO")
    cat.sheet_state = "veryHidden"
    cat.append(["CLAVE", "DESCRIPCION"])
    cat.append(["A", "Catálogo inventado"])
    return _bytes(wb)


def ruta_xlsx() -> bytes:
    """Formato 'Lay Out': fila 1 agrupada y repetida, fila 2 con los nombres."""
    wb = Workbook()
    ws = wb.active
    ws.title = "RUTA"
    ws.append(["DATOS DE CONTACTO", None, None, "DATOS DE DIRECCIÓN", "DATOS DE DIRECCIÓN", "DATOS DE DIRECCIÓN",
               "CASO DE CONTAR CON LOS DATOS", "CASO DE CONTAR CON LOS DATOS"])
    ws.append(["FOLIO", "No SERIE", "MODELO", "CALLE", "COLONIA", "MUNICIPIO", "TELEFONO", "COMENTARIO LIBRE"])
    ws.append(["F-1", "SER00001", "MOD-A", "Calle Uno 1", "Col Uno", "Mpio Uno", "5550100", None])
    ws.append([None] * 8)  # vacía
    ws.append(["Nota al pie de página"] + [None] * 7)  # una celda
    ws.append([None, None, None, "Calle Dos 2", "Col Dos", "Mpio Dos", "5550101", None])  # sin clave
    ws.append(["F-2", None, "MOD-B", "Calle Tres 3", "Col Tres", "Mpio Tres", None, "texto"])
    ws.append(["FOLIO", "No SERIE", "MODELO", "CALLE", "COLONIA", "MUNICIPIO", "TELEFONO", None])  # encabezado repetido
    for nombre in ("LinesByCategori", "ModelsByLines"):
        cat = wb.create_sheet(nombre)
        cat.sheet_state = "hidden"
        cat.append(["LINEA", "MODELO"])  # un solo término del vocabulario: no es encabezado
        cat.append(["Línea 1", "MOD-A"])
    return _bytes(wb)


def encabezado_en_fila(fila: int) -> bytes:
    """Título(s) arriba y el encabezado en la fila indicada (1 a 4)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    for i in range(1, fila):
        ws.append([f"Título inventado {i}"])
    ws.append(["FOLIO", "No SERIE", "MODELO", "CALLE"])
    ws.append(["F-9", "SER00009", "MOD-Z", "Calle Nueve 9"])
    return _bytes(wb)


def penafiel_ods() -> bytes:
    """ODS con una hoja visible de datos, una oculta con el mismo encabezado y un catálogo oculto."""
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.style import Style, TableProperties
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    oculto = Style(name="taOculta", family="table")
    oculto.addElement(TableProperties(display="false"))
    doc.automaticstyles.addElement(oculto)

    def hoja(nombre, filas, oculta=False):
        t = Table(name=nombre, stylename="taOculta") if oculta else Table(name=nombre)
        for fila in filas:
            tr = TableRow()
            for v in fila:
                tc = TableCell(valuetype="string") if v is not None else TableCell()
                if v is not None:
                    tc.addElement(P(text=str(v)))
                tr.addElement(tc)
            t.addElement(tr)
        doc.spreadsheet.addElement(t)

    enc = ["FOLIO", "No SERIE", "MODELO", "CALLE"]
    hoja("DSD", [enc, ["F-1", "SER00001", "MOD-A", "Calle Uno 1"]])
    hoja("DSD_OCULTA", [enc, ["F-2", "SER00002", "MOD-B", "Calle Dos 2"]], oculta=True)
    hoja("CATALOGO", [["LINEA", "DESCRIPCION"], ["L1", "Catálogo inventado"]], oculta=True)
    buf = BytesIO()
    doc.write(buf)
    return buf.getvalue()


ARCHIVOS = {
    "kof.xlsx": kof_xlsx,
    "ruta.xlsx": ruta_xlsx,
    "encabezado_fila1.xlsx": lambda: encabezado_en_fila(1),
    "encabezado_fila4.xlsx": lambda: encabezado_en_fila(4),
    "penafiel.ods": penafiel_ods,
}


def main() -> int:
    for nombre, fn in ARCHIVOS.items():
        (DIR / nombre).write_bytes(fn())
        print("escrito", nombre)
    return 0


if __name__ == "__main__":
    sys.exit(main())
