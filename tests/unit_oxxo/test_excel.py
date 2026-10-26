"""Lectura de los Excel de garantías de Oxxo (.xlsm sintéticos)."""

from io import BytesIO

import pytest
from openpyxl import Workbook

from parsers.oxxo_parse.excel import ExcelIlegible, es_excel, leer_filas
from tests.fixtures.oxxo import build_oxxo


def _bytes(wb):
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_lee_xlsm_con_filas_y_llaves_canonicas():
    nombre = build_oxxo.EXCEL_INIT
    filas = leer_filas(nombre, (build_oxxo.OUT / nombre).read_bytes())
    assert len(filas) == 3
    assert filas[0]["FOLIO"] == "OX-0001"
    assert filas[0]["SERIE"] == "SERIE-EJ-0101"
    assert filas[0]["TELEFONO"] == "5550100001"
    assert filas[2]["FOLIO"] is None


def test_cuartos_frios_con_titulo_antes_del_encabezado():
    filas = leer_filas(build_oxxo.EXCEL_REPLY, (build_oxxo.OUT / build_oxxo.EXCEL_REPLY).read_bytes())
    assert len(filas) == 4
    assert {f["MODELO"] for f in filas} == {"CUARTO-FRIO-EJ"}


def test_ignora_filas_vacias():
    wb = Workbook()
    ws = wb.active
    ws.append(["SERIE", "MODELO"])
    ws.append(["A1", "M1"])
    ws.append([None, None])
    ws.append(["  ", None])
    ws.append(["A2", "M2"])
    assert [f["SERIE"] for f in leer_filas("x.xlsm", _bytes(wb))] == ["A1", "A2"]


def test_usa_la_primera_hoja_visible():
    wb = Workbook()
    oculta = wb.active
    oculta.title = "oculta"
    oculta.append(["SERIE", "MODELO"])
    oculta.append(["OCULTA", "X"])
    oculta.sheet_state = "hidden"
    visible = wb.create_sheet("visible")
    visible.append(["SERIE", "MODELO"])
    visible.append(["VISIBLE", "Y"])
    assert [f["SERIE"] for f in leer_filas("x.xlsm", _bytes(wb))] == ["VISIBLE"]


def test_encabezados_con_acentos_y_alias():
    wb = Workbook()
    ws = wb.active
    ws.append(["Número de caso/folio", "Teléfono", "Diagnóstico y/o falla"])
    ws.append(["OX-9", "555 0100", "Falla"])
    [fila] = leer_filas("x.xlsm", _bytes(wb))
    assert fila == {"FOLIO": "OX-9", "TELEFONO": "555 0100", "FALLA": "Falla"}


@pytest.mark.parametrize("contenido", [b"no es un excel", b""])
def test_archivo_dañado(contenido):
    with pytest.raises(ExcelIlegible):
        leer_filas("x.xlsm", contenido)


def test_sin_encabezado_reconocible():
    wb = Workbook()
    wb.active.append(["a", "b"])
    wb.active.append([1, 2])
    with pytest.raises(ExcelIlegible, match="sin_encabezado"):
        leer_filas("x.xlsm", _bytes(wb))


@pytest.mark.parametrize("nombre,esperado", [
    ("x.xlsm", True), ("X.XLSX", True), ("x.xls", True), ("~$x.xlsx", False), ("x.pdf", False), (None, False),
])
def test_es_excel(nombre, esperado):
    assert es_excel(nombre) is esperado
