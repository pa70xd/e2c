from parsers.core.tablas import unir_tablas as core_unir_tablas
from parsers.core.texto import normalizar_texto as core_normalizar_texto
from parsers.core.texto import eliminar_stopwords as core_eliminar_stopwords
from parsers.core.texto import convertir_html_a_texto as core_convertir_html_a_texto
from config.config_general.stopwords import STOPWORDS
from config.customers.danone.movimientos_danone import MOVEMENTS,CREATE_REQUESTS,IGNORE_REQUESTS
from config.customers.danone.columns_names import (
    COL_FILA_ORIGINAL,
    COL_HASH_ORIGINAL,
)

import hashlib
import re
import unicodedata
from pathlib import Path
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
import pandas as pd
import base64

class DANONE_UTILS:


    @staticmethod
    def analizar_proposito(html,asunto) -> bool:
        # La decisión pasa por el gate del clasificador; con el proveedor
        # por defecto (regex) es exactamente analizar_proposito_regex.
        from classifier import gate_proposito
        return gate_proposito("danone", asunto, html)

    @staticmethod
    def analizar_proposito_regex(html,asunto) -> bool:
        has_request = False
        has_movement = False

        text = DANONE_UTILS.convertir_html_a_texto(html)
        text = (f"{text.lower()} {asunto.lower()}")

        print(f"[TEXT] asunto: {text}")


        if any(re.search(pattern, text) for pattern in IGNORE_REQUESTS):
            return False

        has_request = any(
            re.search(pattern, text)
            for pattern in CREATE_REQUESTS
        )

        # print(f"[PROPOSITO] is request: {has_request}")

        has_movement = any(
            re.search(pattern, text)
            for pattern in MOVEMENTS
        )

        # print(f"[PROPOSITO] is movement: {has_movement}")

        result = has_request or has_movement

        return result

    convertir_html_a_texto = staticmethod(core_convertir_html_a_texto)

    eliminar_stopwords = staticmethod(core_eliminar_stopwords)

    normalizar_texto = staticmethod(core_normalizar_texto)

    unir_tablas = staticmethod(core_unir_tablas)

    @staticmethod
    def sql_list(values):
        return ", ".join(f"'{v}'" for v in values)

    
    def distinct_type_service(self,table,duckConn):

        type_services = [
            row[0]
            for row in duckConn.execute(
                f'''
                SELECT DISTINCT "TIPO DE SERVICIO"
                FROM {table}
                '''
            ).fetchall()
        ]


        return type_services

    @staticmethod
    def safe_filename(subject, sender, received):
        subject = subject.strip()

        subject = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", subject, flags=re.IGNORECASE)
        subject = re.sub(r'[<>:"/\\|?*]', "", subject)
        subject = re.sub(r"\s+", " ", subject).strip()

        sender = sender.replace("<", "").replace(">", "").strip()

        received = received.strftime("%Y-%m-%d %H-%M")

        return f"{subject} - {sender} - {received}"

    @staticmethod
    def mark_rejected_original_files(duck_conn, temporary_files):
        rows = duck_conn.execute(f"""
            SELECT "{COL_HASH_ORIGINAL}", "{COL_FILA_ORIGINAL}", "OS"
            FROM solicitudes_rechazadas
            WHERE "{COL_HASH_ORIGINAL}" IS NOT NULL
        """).fetchall()

        rejected_by_file = {}
        for file_hash, row_number, reason in rows:
            rejected_by_file.setdefault(str(file_hash), {})[
                int(row_number)
            ] = reason

        if not rejected_by_file:
            return

        fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        for _, file_path in temporary_files:
            path = Path(file_path)
            if path.suffix.lower() != ".xlsx":
                print(
                    f"[WARN] No se pueden pintar filas en {path.name}; "
                    "solo se admite formato .xlsx"
                )
                continue

            workbook = None
            try:
                file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
                rejected_rows = rejected_by_file.get(file_hash, {})
                if not rejected_rows:
                    continue

                workbook = load_workbook(path)
                worksheet = next(
                    (
                        sheet
                        for sheet in workbook.worksheets
                        if sheet.sheet_state == "visible"
                    ),
                    workbook.active,
                )
                motivo_column = worksheet.max_column + 1
                worksheet.cell(
                    row=1,
                    column=motivo_column,
                    value="OS",
                )

                for row_number, reason in rejected_rows.items():
                    for cell in worksheet[row_number]:
                        cell.fill = fill
                    if reason:
                        worksheet.cell(
                            row=row_number,
                            column=motivo_column,
                            value=str(reason),
                        )
                        worksheet.cell(
                            row=row_number,
                            column=1,
                        ).comment = Comment(str(reason), "REPARE")

                workbook.save(path)
            except Exception as error:
                print(
                    f"[WARN] No se pudo marcar el Excel original "
                    f"{path.name}: {error}"
                )
            finally:
                if workbook is not None:
                    workbook.close()


