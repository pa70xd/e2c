from config.config_general.stopwords import STOPWORDS
from openpyxl.styles import PatternFill
# from config.customers.penafiel.movimientos_penafiel import MOVEMENTS,CREATE_REQUESTS,IGNORE_REQUESTS
from pathlib import Path
from openpyxl import load_workbook
from parsers.gen_parse.reader import GEN_READER
import re
import unicodedata
from bs4 import BeautifulSoup

gen_reader = GEN_READER()
class GEN_UTILS:


    @staticmethod
    def analizar_proposito( 
            text, 
            MOVEMENTS,
            CREATE_REQUESTS,
            IGNORE_REQUESTS
        ) -> bool:


        has_request = False
        has_movement = False

        # El texto llega en HTML/mayúsculas; los patrones esperan texto plano en minúsculas sin acentos.
        texto_plano = BeautifulSoup(text or "", "html.parser").get_text(separator=" ")
        texto_plano = GEN_UTILS.normalizar_texto(texto_plano)

        print(f"[TEXT] asunto: {texto_plano}")


        if any(re.search(pattern, texto_plano) for pattern in IGNORE_REQUESTS):
            return False

        has_request = any(
            re.search(pattern, texto_plano)
            for pattern in CREATE_REQUESTS
        )

        print(f"[PROPOSITO] is request: {has_request}")

        has_movement = any(
            re.search(pattern, texto_plano)
            for pattern in MOVEMENTS
        )

        print(f"[PROPOSITO] is movement: {has_movement}")

        result = has_request or has_movement

        return result

    @staticmethod
    def convertir_html_a_texto(html: str) -> str:
        texto = BeautifulSoup(html, "html.parser").get_text(separator="\n")
        texto = GEN_UTILS.normalizar_texto(texto)
        texto = GEN_UTILS.eliminar_stopwords(texto)
        return texto

    @staticmethod
    def eliminar_stopwords(texto: str) -> str:
        texto = texto.lower()
        texto = re.sub(r"[^\w\s]", " ", texto)
        palabras = texto.split()
        palabras = [p for p in palabras if p not in STOPWORDS]
        return " ".join(palabras)

    @staticmethod
    def normalizar_texto(texto: str) -> str:
        if not texto:
            return ""

        texto = texto.lower()
        texto = unicodedata.normalize("NFD", texto)
        texto = "".join(
            c for c in texto
            if unicodedata.category(c) != "Mn"
        )
        texto = re.sub(r"[^\w\s]", " ", texto)
        texto = re.sub(r"\s+", " ", texto)

        return texto.strip()

    @staticmethod
    def unir_tablas(tables: list[dict], duck_conn) -> str:
        if not tables:
            return None

        tabla_destino = tables[0]["table"]

        # Obtener columnas de la primera tabla
        columnas = duck_conn.execute(f"""
            SELECT column_name
            FROM information_schema.columns
            WHERE table_name = '{tabla_destino}'
            ORDER BY ordinal_position
        """).fetchall()

        columnas = [c[0] for c in columnas]

        selects = [f'SELECT * FROM "{tabla_destino}"']

        for tabla in tables[1:]:
            nombre = tabla["table"]

            cols = duck_conn.execute(f"""
                SELECT column_name
                FROM information_schema.columns
                WHERE table_name = '{nombre}'
                ORDER BY ordinal_position
            """).fetchall()

            cols = [c[0] for c in cols]

            # Si la estructura cambia, regresar la primera tabla sin modificar
            if cols != columnas:
                return tabla_destino

            selects.append(f'SELECT * FROM "{nombre}"')

        union_sql = "\nUNION ALL\n".join(selects)

        duck_conn.execute(f"""
            CREATE OR REPLACE TABLE "{tabla_destino}" AS
            {union_sql}
        """)

        return tabla_destino

    @staticmethod
    def sql_list(values):
        return ", ".join(f"'{v}'" for v in values)

    
    def distinct_type_service(self,table,duckConn):

        type_services = [
            row[0]
            for row in duckConn.execute(
                f'''
                SELECT DISTINCT "TIPO DE SERVICIO"
                FROM {table}
                '''
            ).fetchall()
        ]


        return type_services

    @staticmethod
    def safe_filename(subject, sender, received):
        subject = subject.strip()

        subject = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", subject, flags=re.IGNORECASE)
        subject = re.sub(r'[<>:"/\\|?*]', "", subject)
        subject = re.sub(r"\s+", " ", subject).strip()

        sender = sender.replace("<", "").replace(">", "").strip()

        received = received.strftime("%Y-%m-%d %H-%M")

        return f"{subject} - {sender} - {received}"

    @staticmethod
    def mark_rejected_original_files(duck_conn, temporary_files,NAMES_COLUMNS):
        rejected_rows = duck_conn.execute(
            'SELECT "SERIE", "OS" FROM solicitudes_rechazadas '
            'WHERE "SERIE" IS NOT NULL'
        ).fetchall()
        rejected_by_serie = {
            str(serie).strip().upper(): reason
            for serie, reason in rejected_rows
            if str(serie).strip()
        }

        if not rejected_by_serie:
            return

        column_names = {
            gen_reader.normalize(name): value
            for name, value in NAMES_COLUMNS.items()
        }
        yellow_fill = PatternFill(fill_type="solid", fgColor="FFFF00")

        for _, file_path in temporary_files:
            path = Path(file_path)
            if path.suffix.lower() not in (".xlsx", ".xlsm"):
                print(
                    f"[WARN] No se pueden marcar rechazos en {path.name}; "
                    "solo se admite formato .xlsx o .xlsm"
                )
                continue

            try:
                workbook = load_workbook(
                    path,
                    keep_vba=path.suffix.lower() == ".xlsm",
                )
                worksheet = next(
                    (
                        sheet
                        for sheet in workbook.worksheets
                        if sheet.sheet_state == "visible"
                    ),
                    workbook.active,
                )

                serie_column = None
                os_column = None
                for cell in worksheet[1]:
                    normalized = gen_reader.normalize(cell.value or "")
                    canonical = column_names.get(normalized, normalized)
                    if canonical == "SERIE":
                        serie_column = cell.column
                    elif canonical == "OS":
                        os_column = cell.column

                if serie_column is None:
                    workbook.close()
                    continue

                if os_column is None:
                    os_column = worksheet.max_column + 1
                    worksheet.cell(row=1, column=os_column, value="OS")

                for row_number in range(2, worksheet.max_row + 1):
                    serie = worksheet.cell(row=row_number, column=serie_column).value
                    if serie is None:
                        continue

                    reason = rejected_by_serie.get(str(serie).strip().upper())
                    if not reason:
                        continue

                    worksheet.cell(row=row_number, column=os_column, value=reason)
                    for cell in worksheet[row_number]:
                        cell.fill = yellow_fill

                workbook.save(path)
                workbook.close()
            except Exception as error:
                print(f"[WARN] No se pudo marcar el Excel original {path.name}: {error}")








