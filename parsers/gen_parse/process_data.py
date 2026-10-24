# from config.customers.penafiel.movimientos_penafiel import MOVEMENTS,CREATE_REQUESTS,IGNORE_REQUESTS
# from config.customers.penafiel.movimientos_penafiel import DIC_SABANAS_MOV
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from models.data_models import EmailModel
from dotenv import load_dotenv
from openpyxl import Workbook
from datetime import datetime
from io import BytesIO
import pandas as pd
import requests
import re
import os

load_dotenv()

class GEN_PROCESS:

    def __init__(self, duckConn):
        self.duckConn = duckConn


    # def detect_purpose(self,text: str) -> bool:
    #     """ Funcion para detectra proposito SOLICITUD valido y descarta solicitudes invalidas """
    #     text = text.lower()

    #     if any(re.search(pattern, text) for pattern in IGNORE_REQUESTS):
    #         return False

    #     has_request = any(
    #         re.search(pattern, text)
    #         for pattern in CREATE_REQUESTS
    #     )

    #     has_movement = any(
    #         re.search(pattern, text)
    #         for pattern in MOVEMENTS
    #     )

    #     return has_request and has_movement


    def fill_email_model(self, email, customer = 'penafiel') -> EmailModel:

        model = EmailModel(
            graph_message_id=email.id,
            conversation_id=email.conversation_id,
            sender_email=(
                email.sender.email_address.address
                if email.sender and email.sender.email_address
                else None
            ),
            subject=email.subject,
            direction="INBOUND",
            customer=customer,
            link_email=email.web_link,
        )

        # Microsoft Graph
        model.internet_message_id = email.internet_message_id
        model.conversation_index = email.conversation_index
        model.parent_message_id = None
        model.change_key = email.change_key

        # Remitente
        if email.sender and email.sender.email_address:
            model.sender_name = email.sender.email_address.name

        # Destinatarios
        model.recipients_to = [
            {
                "name": r.email_address.name,
                "email": r.email_address.address
            }
            for r in (email.to_recipients or [])
            if r.email_address
        ]

        model.recipients_cc = [
            {
                "name": r.email_address.name,
                "email": r.email_address.address
            }
            for r in (email.cc_recipients or [])
            if r.email_address
        ]

        model.recipients_bcc = [
            {
                "name": r.email_address.name,
                "email": r.email_address.address
            }
            for r in (email.bcc_recipients or [])
            if r.email_address
        ]

        # Contenido
        model.body = (
            email.body.content
            if email.body
            else None
        )

        # Adjuntos
        model.has_attachments = email.has_attachments
        model.attachment_count = len(email.attachments or [])

        # Estado
        model.processing_status = "RECEIVED"
        model.processing_message = None

        # Fechas
        model.received_at = email.received_date_time
        model.sent_at = email.sent_date_time

        # Auditoría
        now = datetime.now()
        model.created_at = now
        model.updated_at = now

        return model

    def extract_serie_id(self,table):

        df = self.duckConn.execute(
            f"""
            SELECT DISTINCT
                "ID CLIENTE",
                "SERIE DNN/BNF",
                "SERIE FABRICANTE"
            FROM {table}
            WHERE "ID CLIENTE" IS NOT NULL
            AND TRIM(CAST("ID CLIENTE" AS VARCHAR)) NOT IN ('', 'na', 'nan')
            """
        ).fetchdf()

        id_clientes = (
            df["ID CLIENTE"]
            .dropna()
            .astype(str)
            .tolist()
        )

        series = (
            df["SERIE DNN/BNF"]
            .dropna()
            .astype(str)
            .tolist()
        )

        activos = (
            df["SERIE FABRICANTE"]
            .dropna()
            .astype(str)
            .tolist()
        )


        return id_clientes, series, activos


    def fill_movimiento(self,table):
        values = ", ".join(f"'{v}'" for v in DIC_SABANAS_MOV["MOVIMIENTOS"])


        self.duckConn.execute(f"""
            INSERT INTO sabana_movimiento (
                Distribuidora,
                Categoria,
                Linea,
                Modelo,
                Numero_Activo,
                Serie,        
                Numero_Puertas,
                Imagen,
                Canal,
                Hora_Atencion, 
                Observaciones,
                Tipo_Servicio,
                Id_Servicio,
                Tipo_Actividad,
                Actividad,
                Datos_Prospeccion,
                Tipo_Orden,
                Numero_Cliente,
                Nombre_PDV,
                Calle,
                Numero,
                Entre_Calle1,
                Entre_Calle2,
                Colonia,
                Municipio,
                Codigo_Postal,
                Telefono_Fax,
                Nombre_Contacto_PDV,
                Correo_Electronico,
                Latitud,
                Longitud,
                Distribuidora_Puesto_Trabajo
            )
            SELECT
                COALESCE("DISTRIBUIDORA_GEN",'') as Distribuidora,                "CATEGORIA_GEN" as Categoria,
                "LINEA_GEN" as Linea,
                "MODELO_GEN" as Modelo,
                "ACTIVO_GEN" as Numero_Activo,
                COALESCE("SERIE_GEN","SERIE DNN/BNF") as Serie,
                '' as Numero_Puertas,
                '' as Imagen,
                'SC | Sin Canal' as Canal,
                '' as Hora_Atencion,

                COALESCE("TIPO DE SERVICIO", '') || ' / ' ||
                COALESCE(COMENTARIOS, '') || ' /Region: ' ||
                COALESCE(REGION, '') || ' /Canal: ' ||
                COALESCE(CANAL, '') || ' /Ruta: ' ||
                COALESCE(RUTA, '') AS Observaciones,

                '' as Tipo_Servicio,
                '' as Id_Servicio,
                '' as Tipo_Actividad,
                '' as Actividad,

                'NA' as Datos_Prospeccion,
                '' as Tipo_Orden,
                "ID CLIENTE" as Numero_Cliente,
                "RAZON SOCIAL" as Nombre_PDV,
                "DIRECCION" as Calle,
                'SN' as Numero,
                'NP' as Entre_Calle1,
                'NP' as Entre_Calle2,
                "COLONIA" as Colonia,
                CIUDAD as Municipio,
                CP as Codigo_Postal,
                "TELEFONO 2" as Telefono_Fax,
                '' as Nombre_Contacto_PDV,
                '' as Correo_Electronico,
                LATITUD as Latitud,
                LONGITUD as Longitud,
                "DISTRIBUIDORA_GEN" as Distribuidora_Puesto_Trabajo
            FROM {table}

            WHERE

                "TIPO DE SERVICIO"  in ({values})
        """)



    def fill_ruta(self,table):

        values = ", ".join(f"'{v}'" for v in DIC_SABANAS_MOV["RUTA"])

        self.duckConn.execute(f"""
            INSERT INTO sabana_ruta (
                Distribuidora,
                Categoria,
                Linea,
                Modelo,
                Numero_Activo,
                Serie,
                Observaciones,
                Hora_Atencion,
                Prioridad,
                Folio,
                Ticket_Solicitud,
                Canal,
                Correo_Electronico_1,
                Numero_Cliente,
                Nombre_PDV,
                Calle,
                Numero,
                Entre_Calle1,
                Entre_Calle2,
                Colonia,
                Municipio,
                Codigo_Postal,
                Latitud,
                Longitud,
                Telefono_Fax_1,
                Telefono_Fax_2,
                Correo_Electronico_2,
                Nombre_Contacto_PDV
            )
            SELECT
                COALESCE("DISTRIBUIDORA_GEN",'') as Distribuidora,
                "CATEGORIA_GEN" as Categoria,
                "LINEA_GEN" as Linea,
                "MODELO_GEN" as Modelo,
                "ACTIVO_GEN" as Numero_Activo,
                COALESCE("SERIE_GEN","SERIE DNN/BNF") as Serie,

                COALESCE("TIPO DE SERVICIO", '') || ' / ' ||
                COALESCE(COMENTARIOS, '') || ' /Region: ' ||
                COALESCE(REGION, '') || ' /Canal: ' ||
                COALESCE(CANAL, '') || ' /Ruta: ' ||
                COALESCE(RUTA, '') AS Observaciones,

                '' as Hora_Atencion,
                '' as Prioridad,
                '' as Folio,
                '' as Ticket_Solicitud,
                '' as Canal,
                '' as Correo_Electronico_1,
                "TELEFONO 1" as Numero_Cliente,
                "RAZON SOCIAL" as Nombre_PDV,
                DIRECCION as Calle,
                '' as Numero,
                '' as Entre_Calle1,
                '' as Entre_Calle2,
                COLONIA as Colonia,
                CIUDAD as Municipio,
                CP as Codigo_Postal,
                LATITUD as Latitud,
                LONGITUD as Longitud,
                "TELEFONO 2" as Telefono_Fax_1,
                '' as Telefono_Fax_2,
                '' as Correo_Electronico_2,
                '' as Nombre_Contacto_PDV


            FROM {table}

            WHERE

                "TIPO DE SERVICIO"  in ({values})
        """)

    def distinct_type_service(self,tables):

        if tables:

            table = tables[0]["table"]

            type_services = [
                row[0]
                for row in self.duckConn.execute(
                    f'''
                    SELECT DISTINCT "TIPO DE SERVICIO"
                    FROM {table}
                    '''
                ).fetchall()
            ]

            return type_services
        return None


    def get_categoria_linea(self, modelo, catalog):
        if pd.isna(modelo):
            return None, None

        modelo = str(modelo).strip().upper()

        for categoria, lineas in catalog.items():
            for linea, modelos in lineas.items():
                for m in modelos:
                    if modelo == str(m).strip().upper():
                        return categoria, linea

        return None, None



    @staticmethod
    def sql_list(values):
        return ", ".join(f"'{v}'" for v in values)


    def union_cap_tables(self,table):

        self.duckConn.execute(f"""
                CREATE OR REPLACE TABLE "{table}" AS
                SELECT
                DISTINCT
                    s.*,
                    -- Información Punto de Venta
                    p.DISTRIBUIDORA_GEN AS DISTRIBUIDORA_GEN,
                    p.CLIENTE as CLIENTE,

                    -- Información Equipo
                    m.ACTIVO      AS ACTIVO_GEN,
                    LTRIM(m.SERIE, '0') AS SERIE_GEN,
                    COALESCE(m.MODELO, s.MODELO) AS MODELO_GEN,
                    m.CATEGORIA   AS CATEGORIA_GEN,
                    m.LINEA       AS LINEA_GEN

                FROM "{table}" AS s

                LEFT JOIN tmp_mpto AS p
                    ON CAST(s."ID CLIENTE" AS VARCHAR) = p.ID_PTO_VENTA

                LEFT JOIN tmp_modl AS m
                    ON
                        CAST(s."SERIE DNN/BNF" AS VARCHAR) = m.ACTIVO

                    OR

                        LTRIM(CAST(s."SERIE FABRICANTE" AS VARCHAR), '0')
                            =
                        LTRIM(m.SERIE, '0')
            """)



    def complet_catalog_line(self, table: str, catalog: dict):

        df = self.duckConn.execute(f"""
            SELECT rowid, MODELO_GEN
            FROM {table}
            WHERE CATEGORIA_GEN IS NULL
            OR CATEGORIA_GEN = ''
            OR LINEA_GEN IS NULL
            OR LINEA_GEN = ''
        """).fetchdf()

        for _, row in df.iterrows():

            modelo = row["MODELO_GEN"]

            if pd.isna(modelo):
                continue

            modelo = str(modelo).strip().upper().split()[0]

            categoria, linea = self.get_categoria_linea(modelo, catalog)

            if categoria is not None:
                self.duckConn.execute("""
                    UPDATE {table}
                    SET
                        CATEGORIA_GEN = ?,
                        LINEA_GEN = ?
                    WHERE rowid = ?
                """.format(table=table), [categoria, linea, row["rowid"]])


    def prepare_data_request(
        self,
        id_email,
        customer='penafiel',
        mail = None,
    ):

        total_coolers = self.duckConn.execute("""
            SELECT COUNT(*)
            FROM solicitud
        """).fetchone()[0]

        # failed_coolers = self.duckConn.execute("""
        #     SELECT COUNT(*)
        #     FROM solicitudes_rechazadas
        # """).fetchone()[0]

       
        return {
            "request_code": id_email,
            "customer": customer,
            "status": "IN_PROGRESS",
            "priority": "HIGH",
            "total_coolers": total_coolers,
            "pending_coolers": total_coolers,
            "completed_coolers": 0,
            "failed_coolers": 0,
            "assigned_to": "NA",
        }



    def prepare_data_cooler(
        self,
        id_email,
    ):
        df = self.duckConn.execute("""
            SELECT *
            FROM solicitud
        """).fetchdf()

        def clean_value(value):
            if value is None:
                return None

            if pd.isna(value):
                return None

            if isinstance(value, str):
                value = value.strip()

                if value.lower() in ("nan", "none", ""):
                    return None

                return value

            return value

        def clean_serie(value):
            value = clean_value(value)

            if value is None:
                return None

            return str(value).strip()

        detalles = []

        for i, row in df.iterrows():

            detalles.append({
                # ==================================================
                # IDENTIFICACIÓN
                # ==================================================
                "request_code": id_email,
                "row_number": i + 1,

                # ==================================================
                # DATOS DE GESTIÓN
                # ==================================================
                "status": "PENDING",
                "rejection_reason": None,
                "prioridad": None,
                "folio": None,
                "ticket_solicitud": None,
                "assigned_to": None,
                "gestion_observaciones": None,
                "fecha_asignacion": None,
                "fecha_atencion": None,
                "fecha_completado": None,

                # ==================================================
                # DATOS ORIGINALES DE LA SOLICITUD
                # ==================================================
                "id_cliente": clean_value(row.get("ID CLIENTE")),
                "fecha_solicitud": clean_value(row.get("FECHA DE SOLICITUD")),
                "modelo": clean_value(row.get("MODELO")),

                # SERIES  SI VIENE NaN SE GUARDA COMO None
                "serie_fabricante": clean_serie(
                    row.get("SERIE FABRICANTE")
                ),
                "serie_dnn_bnf": clean_serie(
                    row.get("SERIE DNN/BNF")
                ) or clean_serie(row.get("SERIE")),

                "tipo_servicio": clean_value(row.get("TIPO DE SERVICIO")),
                "canal": clean_value(row.get("CANAL")),
                "region": clean_value(row.get("REGION")),
                "plaza": clean_value(row.get("PLAZA")),
                "razon_social": clean_value(row.get("RAZON SOCIAL")),
                "direccion": clean_value(row.get("DIRECCION")),
                "cp": clean_value(row.get("CP")),
                "ciudad": clean_value(row.get("CIUDAD")),
                "colonia": clean_value(row.get("COLONIA")),
                "latitud": clean_value(row.get("LATITUD")),
                "longitud": clean_value(row.get("LONGITUD")),
                "ruta": clean_value(row.get("RUTA")),
                "telefono_1": clean_value(row.get("TELEFONO 1")),
                "telefono_2": clean_value(row.get("TELEFONO 2")),
                "comentarios": clean_value(row.get("COMENTARIOS")),
                "cbu": clean_value(row.get("CBU")),
                "cliente": clean_value(row.get("CLIENTE")),
                "os": clean_value(row.get("OS")),
            })

        return detalles


    def get_address(self,latitud, longitud):

        url = "https://maps.googleapis.com/maps/api/geocode/json"
        API_KEY = os.getenv('MAPS_KEY')


        params = {
            "latlng": f"{latitud},{longitud}",
            "key": API_KEY,
            "language": "es",
        }

        response = requests.get(url, params=params)

        if response.status_code != 200:
            raise Exception(f"Error HTTP: {response.status_code}")

        data = response.json()

        if data["status"] != "OK":
            raise Exception(
                f"Error de Google Maps: {data['status']} - "
                f"{data.get('error_message', '')}"
            )

        resultado = data["results"][0]

        direccion = resultado.get("formatted_address")

        datos = {
            "direccion": direccion,
            "codigo_postal": None,
            "colonia": None,
            "ciudad": None,
            "estado": None,
            "pais": None,
        }

        for componente in resultado["address_components"]:
            tipos = componente["types"]
            nombre = componente["long_name"]

            if "postal_code" in tipos:
                datos["codigo_postal"] = nombre

            elif "sublocality" in tipos or "neighborhood" in tipos:
                datos["colonia"] = nombre

            elif "locality" in tipos:
                datos["ciudad"] = nombre

            elif "administrative_area_level_1" in tipos:
                datos["estado"] = nombre

            elif "country" in tipos:
                datos["pais"] = nombre

        return datos


    def prepare_cooler_update(self, table_name: str):

        columnas = {
            fila[0]
            for fila in self.duckConn.execute(
                f'DESCRIBE SELECT * FROM {table_name}'
            ).fetchall()
        }

        series_dnn = []
        if "SERIE DNN/BNF" in columnas:
            series_dnn.append(
                'NULLIF(TRIM(CAST("SERIE DNN/BNF" AS VARCHAR)), \'\')'
            )
        if "SERIE" in columnas:
            series_dnn.append(
                'NULLIF(TRIM(CAST("SERIE" AS VARCHAR)), \'\')'
            )

        serial_dnn = (
            f"COALESCE({', '.join(series_dnn)}, '')"
            if series_dnn
            else "''"
        )

        query = f"""
            SELECT
                "ID CLIENTE" AS customer_id,

                COALESCE(TRIM(CAST("SERIE FABRICANTE" AS VARCHAR)), '') AS serial_number,

                {serial_dnn} AS serial_dnn,

                COALESCE(TRIM(CAST("OS" AS VARCHAR)), '') AS os,

                CASE
                    WHEN COALESCE(TRIM(CAST("OS" AS VARCHAR)), '') = ''
                        THEN 'REJECTED'

                    WHEN UPPER(CAST("OS" AS VARCHAR)) LIKE '%[RECHAZADO]%'
                        THEN 'PENDING'

                    WHEN LTRIM(CAST("OS" AS VARCHAR), '0') NOT LIKE '403%'
                        THEN 'PENDING'

                    ELSE 'ACCEPTED'
                END AS status

            FROM {table_name};
        """
        df = self.duckConn.execute(query).fetchdf()

        df = df.where(pd.notna(df), None)

        return df


    @staticmethod
    def process_email_request_global(rows):

        if not rows:
            return None

        wb = Workbook()
        ws = wb.active
        ws.title = "Solicitud"

        # RealDictRow -> dict
        rows = [dict(row) for row in rows]

        columns = list(rows[0].keys())

        # Encabezados
        ws.append(columns)

        # Campos que deben ser datetime
        date_columns = {
            "created_at",
            "updated_at",
            "completed_at",
            "fecha_solicitud",
            "fecha_asignacion",
            "fecha_atencion",
            "fecha_completado",
            "cooler_created_at",
            "cooler_updated_at",
        }

        # Datos
        for row in rows:

            values = []

            for column in columns:

                value = row.get(column)

                # Convertir fechas a datetime real
                if value is not None and column in date_columns:

                    if not isinstance(value, datetime):
                        try:
                            value = datetime.fromisoformat(
                                str(value).replace("Z", "+00:00")
                            )
                        except (ValueError, TypeError):
                            pass

                values.append(value)

            ws.append(values)

        # Formato de fechas
        for column_index, column_name in enumerate(columns, start=1):

            if column_name in date_columns:

                for row_index in range(2, ws.max_row + 1):

                    cell = ws.cell(
                        row=row_index,
                        column=column_index
                    )

                    if cell.value is not None:
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"

        # Crear tabla
        if ws.max_row >= 2:

            last_column = get_column_letter(ws.max_column)

            table = Table(
                displayName="SolicitudGlobal",
                ref=f"A1:{last_column}{ws.max_row}"
            )

            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            table.tableStyleInfo = style

            ws.add_table(table)

        # Ajustar columnas
        for column_index in range(1, ws.max_column + 1):

            column_letter = get_column_letter(column_index)

            max_length = 0

            for cell in ws[column_letter]:

                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            ws.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                40
            )

        # Congelar encabezado
        ws.freeze_panes = "A2"

        # Generar archivo en memoria
        output = BytesIO()

        wb.save(output)

        output.seek(0)

        fecha = datetime.now().strftime("%Y-%m-%d")

        output.seek(0)

        return (
            f"estatus_global_{fecha}.xlsx",
            output.getvalue()
        )