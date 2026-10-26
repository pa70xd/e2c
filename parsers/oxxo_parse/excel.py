"""Lectura de los Excel de garantías de Oxxo (Lay Out Repare, ``.xlsx`` o
``.xlsm``; puede ser una línea de cuartos fríos de Imbera).

PENDIENTE (lector-excel): cuando ``parsers/core/lector.py`` llegue a la base,
``leer_filas`` debe delegar en ese lector común y este módulo se reduce al
conteo y al mapeo a ``cooler_request``. Mientras tanto lee la primera hoja
visible con openpyxl, ubica la fila de encabezados por alias
(``config/customers/oxxo/columns_names.py``) y devuelve una fila por renglón
con algún dato, con las llaves canónicas.

Los valores nunca se imprimen."""

from __future__ import annotations

import re
import unicodedata
from io import BytesIO

from openpyxl import load_workbook

from config.customers.oxxo.columns_names import (
    FILAS_BUSQUEDA_ENCABEZADO,
    MAX_FILAS_EXCEL,
    MIN_ALIAS_ENCABEZADO,
    NAMES_EXCEL,
)

EXTENSIONES_EXCEL = (".xlsx", ".xlsm")


class ExcelIlegible(Exception):
    """El archivo no se pudo abrir o no tiene un encabezado reconocible."""


def es_excel(nombre: str | None) -> bool:
    return (nombre or "").lower().endswith(EXTENSIONES_EXCEL)


def _encabezado(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(c for c in t if not unicodedata.combining(c)).upper()
    return re.sub(r"\s+", " ", t).strip(" :")


def _celda(valor) -> str | None:
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def leer_filas(contenido: bytes) -> list[dict[str, str | None]]:
    """Filas del Excel con llaves canónicas (NAMES_EXCEL). Una fila cuenta
    si tiene al menos un valor. ``ExcelIlegible`` si el archivo no abre, no
    tiene encabezado reconocible o tiene demasiadas filas."""
    try:
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except Exception as e:  # zip dañado, no es Excel, etc.
        raise ExcelIlegible(type(e).__name__) from e

    try:
        hoja = next((h for h in libro.worksheets if h.sheet_state == "visible"), None)
        if hoja is None:
            raise ExcelIlegible("sin_hoja_visible")
        renglones = list(hoja.iter_rows(values_only=True, max_row=MAX_FILAS_EXCEL + FILAS_BUSQUEDA_ENCABEZADO + 1))
    finally:
        libro.close()

    posicion = None
    for i, renglon in enumerate(renglones[:FILAS_BUSQUEDA_ENCABEZADO]):
        canonicos = [NAMES_EXCEL.get(_encabezado(v)) for v in renglon if v is not None]
        if len({c for c in canonicos if c}) >= MIN_ALIAS_ENCABEZADO:
            posicion = i
            break
    if posicion is None:
        raise ExcelIlegible("sin_encabezado")

    columnas = []
    for v in renglones[posicion]:
        nombre = NAMES_EXCEL.get(_encabezado(v)) if v is not None else None
        columnas.append(nombre if nombre not in columnas else None)

    filas = []
    for renglon in renglones[posicion + 1:]:
        fila = {c: _celda(v) for c, v in zip(columnas, renglon) if c}
        if any(fila.values()):
            filas.append(fila)
    if len(filas) > MAX_FILAS_EXCEL:
        raise ExcelIlegible("demasiadas_filas")
    return filas
