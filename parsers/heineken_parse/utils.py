from io import BytesIO
import re
import unicodedata

from kiota_abstractions.base_request_configuration import RequestConfiguration
from msgraph.generated.users.item.messages.item.message_item_request_builder import MessageItemRequestBuilder
from openpyxl import load_workbook

from config.customers.heineken.columns_names import COL_ARCHIVO, COL_FILA_EXCEL
from config.customers.heineken.reglas_heineken import (
    ASUNTO_D2D, ASUNTO_RUTA, FLUJO_D2D, FLUJO_RUTA,
)
from parsers.holanda_parse.utils import HOLANDA_UTILS

PATRON_FOLIO = re.compile(r"REPARE_\d+", re.IGNORECASE)
_PREFIJOS_RESPUESTA =re.compile(r"^\s*((re|rv|fw|fwd)\s*:\s*)+", re.IGNORECASE)
_ETIQUETAS = re.compile(r"\[[^\]]*\]")


def normalizar_asunto(asunto: str | None) -> str:
    texto = _ETIQUETAS.sub(" ", asunto or "")
    texto = _PREFIJOS_RESPUESTA.sub("", texto)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.split()).upper()


def flujo_por_asunto(asunto: str | None) -> str | None:
    """FLUJO_D2D, FLUJO_RUTA o None si el correo no se procesa."""
    texto = normalizar_asunto(asunto)
    if all(p in texto for p in ASUNTO_D2D):
        return FLUJO_D2D
    if all(p in texto for p in ASUNTO_RUTA):
        return FLUJO_RUTA
    return None


async def resolver_folio(client, user_id, mail_id, mail) -> str | None:
    """Folio REPARE_xxx del texto nuevo o, si no, del cuerpo completo (el correo del bot citado)."""
    encontrado = PATRON_FOLIO.search(getattr(getattr(mail, "unique_body", None), "content", "") or "")
    if encontrado:
        return encontrado.group(0).upper()

    try:
        query_params = MessageItemRequestBuilder.MessageItemRequestBuilderGetQueryParameters(select=["body"])
        config = RequestConfiguration(query_parameters=query_params)
        config.headers.add("Prefer", 'outlook.body-content-type="text"')
        mensaje = await (
            client.users.by_user_id(user_id).messages.by_message_id(mail_id)
            .get(request_configuration=config)
        )
        cuerpo = getattr(getattr(mensaje, "body", None), "content", "") or ""
    except Exception as e:
        print(f"[HEINEKEN][BOT][ERROR] No se pudo leer el cuerpo completo del correo: {e}")
        return None

    encontrado = PATRON_FOLIO.search(cuerpo)
    return encontrado.group(0).upper() if encontrado else None


def escribir_series(content: bytes, series: dict[int, str], col_serie: int | None) -> bytes:
    """Escribe en la columna SERIE del original las series genéricas generadas."""
    if not series or not col_serie:
        return content
    try:
        wb = load_workbook(BytesIO(content))
        ws = wb.worksheets[0]
        for fila, serie in series.items():
            ws.cell(row=fila, column=col_serie).value = serie
        salida = BytesIO()
        wb.save(salida)
        return salida.getvalue()
    except Exception as e:
        print(f"[HEINEKEN][WARN] No se pudieron escribir las series en el original: {e}")
        return content


def originales_pintados(duck_conn, archivos: list[dict],
                        series: dict[str, dict[int, str]] | None = None) -> list[tuple[str, bytes]]:
    """Archivos recibidos con las series genéricas escritas y las rechazadas pintadas."""
    rechazadas: dict[str, dict[int, str]] = {}
    for archivo, fila, motivo in duck_conn.execute(
        f'SELECT "{COL_ARCHIVO}", "{COL_FILA_EXCEL}", MOTIVO FROM solicitudes_rechazadas'
    ).fetchall():
        if archivo and fila:
            rechazadas.setdefault(archivo, {})[int(fila)] = motivo

    series = series or {}
    resultado = []
    for a in archivos:
        if not a.get("content"):
            continue
        content = escribir_series(a["content"], series.get(a["table"], {}), a.get("col_serie"))
        content = HOLANDA_UTILS.pintar_rechazadas(
            content, rechazadas.get(a["table"], {}), a.get("fila_encabezado", 1)
        )
        resultado.append((a["name"], content))
    return resultado
