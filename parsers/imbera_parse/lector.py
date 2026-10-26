"""Adaptador al lector común de Excel (Imbera y Torrey).

El lector común ``parsers/core/lector.py`` lo escribe la sesión
lector-excel. Mientras no esté en la base, ``leer_hojas`` usa openpyxl
directo. Cuando llegue, solo cambia ``leer_hojas`` (mismo contrato: una
``Hoja`` por pestaña, con su estado de visibilidad, filas crudas y celdas
combinadas); el resto del parser no se toca.

TODO(lector-excel): reemplazar ``_leer_con_openpyxl`` por el lector común.
"""

from __future__ import annotations

import datetime as dt
import unicodedata
from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import load_workbook


@dataclass
class Hoja:
    nombre: str
    visible: bool
    # Valores crudos por fila (fila 1 de Excel = índice 0).
    filas: list[list]
    # Rangos combinados como (fila_min, col_min, fila_max, col_max), 1-indexados.
    combinadas: list[tuple[int, int, int, int]] = field(default_factory=list)


def _leer_con_openpyxl(content: bytes) -> list[Hoja]:
    wb = load_workbook(BytesIO(content), data_only=True)
    try:
        hojas = []
        for ws in wb.worksheets:
            hojas.append(Hoja(
                nombre=ws.title,
                visible=ws.sheet_state == "visible",
                filas=[list(fila) for fila in ws.iter_rows(values_only=True)],
                combinadas=[
                    (r.min_row, r.min_col, r.max_row, r.max_col)
                    for r in ws.merged_cells.ranges
                ],
            ))
        return hojas
    finally:
        wb.close()


leer_hojas = _leer_con_openpyxl


# ---------------------------------------------------------------------------
# Normalización de celdas
# ---------------------------------------------------------------------------

def normalizar(texto) -> str:
    """Mayúsculas, sin acentos y con espacios simples ("Ñ" se conserva para
    no confundir "AÑO" con "ANO"; "N°" se conserva)."""
    if texto is None:
        return ""
    texto = str(texto).replace("ñ", "\0").replace("Ñ", "\0")
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = texto.replace("\0", "Ñ").upper()
    return " ".join(texto.split())


def celda_a_texto(valor) -> str:
    """Valor de celda como texto estable: fechas en ISO, enteros sin ".0"."""
    if valor is None:
        return ""
    if isinstance(valor, dt.datetime):
        if valor.time() == dt.time(0, 0):
            return valor.date().isoformat()
        return valor.isoformat(sep=" ")
    if isinstance(valor, dt.date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()
