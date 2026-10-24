from copy import copy
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from parsers.danone_parse.utils import DANONE_UTILS

# Relleno para las filas rechazadas en el archivo original que se regresa.
RELLENO_RECHAZADA = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


class HOLANDA_UTILS:
    """Utilerías genéricas de Danone, reusadas tal cual, más el pintado del
    archivo original."""

    unir_tablas = staticmethod(DANONE_UTILS.unir_tablas)

    @staticmethod
    def pintar_rechazadas(content: bytes, rechazadas: dict[int, str], fila_encabezado: int = 1) -> bytes:
        """Regresa `content` (el Excel original) con las filas de
        `rechazadas` ({fila_excel: motivo}) pintadas en amarillo y el motivo
        en la columna OS (si no existe, se agrega al final). Solo se toca la
        primera hoja. Si no se puede abrir (p. ej. .xls), se regresa sin cambios."""

        if not rechazadas:
            return content

        try:
            wb = load_workbook(BytesIO(content))
            ws = wb.worksheets[0]

            encabezados = {
                c.column: str(c.value).strip().upper()
                for c in ws[fila_encabezado] if c.value is not None and str(c.value).strip()
            }
            col_os = next((col for col, v in encabezados.items() if v == "OS"), None)
            if col_os is None:
                ultima = max(encabezados, default=0)
                col_os = ultima + 1
                nueva = ws.cell(row=fila_encabezado, column=col_os, value="OS")
                if ultima:
                    nueva._style = copy(ws.cell(row=fila_encabezado, column=ultima)._style)

            ultima_columna = max(ws.max_column, col_os)
            for fila, motivo in rechazadas.items():
                for columna in range(1, ultima_columna + 1):
                    ws.cell(row=fila, column=columna).fill = RELLENO_RECHAZADA
                ws.cell(row=fila, column=col_os).value = motivo

            salida = BytesIO()
            wb.save(salida)
            return salida.getvalue()

        except Exception as e:
            print(f"[HOLANDA][WARN] No se pudo pintar el archivo original: {e}")
            return content
