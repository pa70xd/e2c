from pathlib import Path
import pandas as pd
from io import BytesIO
import pandas as pd
from google.cloud import storage
from google.oauth2 import service_account
import re
from dotenv import load_dotenv
import os
from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
import zipfile
import re
import xml.etree.ElementTree as ET
from io import BytesIO
import base64

load_dotenv()

PROJECT_ID = os.getenv("PROJECT_ID")
SERVICE_ACCOUNT_FILE = r"C:\keys\imberalink-238816-1694cd22ed85.json"
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"


class  Excel_Ingestion:
    def get_sheet_path(self, zip_bytes, sheet_name):
        with zipfile.ZipFile(BytesIO(zip_bytes)) as z:
            wb_xml = z.read("xl/workbook.xml")
            rels_xml = z.read("xl/_rels/workbook.xml.rels")
            names = set(z.namelist())

        wb_root = ET.fromstring(wb_xml)
        rid = None
        for sheet in wb_root.iter(f"{{{NS_MAIN}}}sheet"):
            if sheet.get("name") == sheet_name:
                rid = sheet.get(f"{{{NS_R}}}id")
                break

        if rid is None:
            raise ValueError(f"Hoja '{sheet_name}' no encontrada en workbook.xml")

        rels_root = ET.fromstring(rels_xml)
        target = None
        for rel in rels_root.iter(f"{{{NS_RELS}}}Relationship"):
            if rel.get("Id") == rid:
                target = rel.get("Target")
                break

        if target is None:
            raise ValueError(f"Relationship '{rid}' no encontrado en workbook.xml.rels")

        target = target.lstrip("/")
        if target.startswith("xl/"):
            candidate = target
        else:
            candidate = "xl/" + target

        if candidate in names:
            return candidate

        base = target.split("/")[-1]
        for n in names:
            if n.endswith("/" + base) or n == base:
                return n

        raise ValueError(f"No se encontró la ruta de la hoja para Target='{target}'. Namelist: {sorted(names)}")

    def restore_x14_validations(self, template_bytes, output_bytes, sheet_name):
        orig_path = self.get_sheet_path(template_bytes, sheet_name)
        with zipfile.ZipFile(BytesIO(template_bytes)) as z:
            orig_xml = z.read(orig_path).decode("utf-8")

        ext_matches = re.findall(r"<extLst>.*?</extLst>", orig_xml, re.DOTALL)
        if not ext_matches:
            return output_bytes
        ext_block = ext_matches[0]

        ext_block = re.sub(r'\s+xr:uid="[^"]*"', "", ext_block)

        new_path = self.get_sheet_path(output_bytes, sheet_name)
        with zipfile.ZipFile(BytesIO(output_bytes)) as zin:
            new_xml = zin.read(new_path).decode("utf-8")

        if "<extLst>" in new_xml:
            new_xml, n = re.subn(r"<extLst>.*?</extLst>", ext_block, new_xml, count=1, flags=re.DOTALL)
            assert n == 1
        else:
            new_xml = new_xml.replace("</worksheet>", ext_block + "</worksheet>")

        ET.fromstring(new_xml)

        out = BytesIO()
        with zipfile.ZipFile(BytesIO(output_bytes)) as zin, \
                zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = new_xml.encode("utf-8") if item.filename == new_path else zin.read(item.filename)
                zout.writestr(item, data)

        result = out.getvalue()
        with zipfile.ZipFile(BytesIO(result)) as check:
            bad = check.testzip()
            assert bad is None, f"Zip corrupto en: {bad}"

        return result

    def extract_excel(self, duck_conn, table_name = 'sabana'):
        debug_dir = Path("debug")
        debug_dir.mkdir(exist_ok=True)

        output_file = debug_dir / f"{table_name}.xlsx"

        df = duck_conn.execute(f"SELECT * FROM {table_name}").fetchdf()
        df.to_excel(output_file, index=False)

        return output_file
    


    def extract_and_load(self, duck_conn, file_name, customer, table_name="sabana"):


        try:

            df = duck_conn.execute(f"SELECT * FROM {table_name}").fetchdf()

            excel_buffer = BytesIO()

            with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name=table_name)

            excel_buffer.seek(0)

            credentials = service_account.Credentials.from_service_account_file(
                SERVICE_ACCOUNT_FILE
            )


            storage_client = storage.Client(
                project=PROJECT_ID,
                credentials=credentials
            )

            bucket_name = os.getenv("REPARE_BUCKET")

            bucket = storage_client.bucket(bucket_name)

            object_name = f"{customer.upper()}/{file_name}.xlsx"

            # blob = bucket.blob(object_name)

            # blob.upload_from_file(
            #     excel_buffer,
            #     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            # )


        except Exception as e:
            print("ERROR DURANTE LA SUBIDA")
            print(type(e).__name__)
            print(e)
            raise






    def get_excel_attachment(
        self,
        duck_conn,
        table_name: str,
        filename: str | None = None,
    ) -> tuple[str, bytes]:

        df = duck_conn.execute(
            f"SELECT * FROM {table_name}"
        ).fetchdf()

        excel_buffer = BytesIO()

        with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
            df.to_excel(
                writer,
                index=False,
                sheet_name=table_name
            )

        excel_buffer.seek(0)

        if filename is None:
            filename = table_name

        return f"{filename}.xlsx", excel_buffer.read()

    def get_excelv2_attachment(
        self,
        duck_conn,
        table_name: str,
        filename: str | None = None,
    ) -> tuple[str, bytes] | None:


        # Validar que la tabla exista
        exists = duck_conn.execute(f"""
            SELECT COUNT(*)
            FROM information_schema.tables
            WHERE table_name = '{table_name}'
        """).fetchone()[0]

        if not exists:
            return None

        df = duck_conn.execute(
            f"SELECT * FROM {table_name}"
        ).fetchdf()

        if df.empty:
            return None

        template_path = {
            'sabana_movimiento': './excel_templates/template_movimientos.xlsx',
            'sabana_ruta': './excel_templates/template_rutas.xlsx',
            'solicitudes_rechazadas': './excel_templates/template_rechazadas.xlsx',
            'pre_sabana': './excel_templates/template_movimientos.xlsx',
            'solicitud': './excel_templates/solicitud.xlsx',
        }.get(table_name, './excel_templates/template_movimientos.xlsx')

        with open(template_path, "rb") as f:
            template_bytes = f.read()

        wb = load_workbook(template_path)
        ws = wb.worksheets[0]

        

        for i, row in enumerate(df.itertuples(index=False)):
            for j, value in enumerate(row):
                ws.cell(row=i + 2, column=j + 1, value=value)

        output = BytesIO()
        wb.save(output)
        output_bytes = output.getvalue()

        output_bytes = self.restore_x14_validations(template_bytes, output_bytes, ws.title)

        if filename is None:
            filename = f"{table_name.upper()}"
        else:
            filename = f"{filename.upper()}|{table_name.upper()}"
        # with open(f"./debug/TEST_{filename}.xlsx", "wb") as f:
        #     f.write(output_bytes)

        return f"{filename}.xlsx", output_bytes